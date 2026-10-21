"""Enterprise report engine for completed EvalForge evaluations.

The engine consumes persisted evaluation payloads only. It does not execute
scenarios, adapters, authentication, RAG, or scoring.
"""

from __future__ import annotations

import html
import json
import os
import re
import textwrap
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any


CLIENT_REPORT_VERSION = "1.0"
EVALFORGE_VERSION = os.getenv("EVALFORGE_VERSION", "1.0")
PDF_TABLE_CELL_MAX_CHARS = 300
PDF_LONG_TEXT_CHUNK_CHARS = 1400
PDF_CODE_BLOCK_CHUNK_CHARS = 2200


@dataclass(frozen=True)
class ReportTheme:
    brand: str = "EvalForge"
    title: str = "AI Agent Assessment Report"
    primary: str = "#1d4ed8"
    ink: str = "#172033"
    muted: str = "#64748b"
    success: str = "#15803d"
    danger: str = "#b91c1c"
    warning: str = "#c2410c"


class EvaluationReportEngine:
    """Build professional report data and downloadable artifacts."""

    def __init__(self, theme: ReportTheme | None = None) -> None:
        self.theme = theme or ReportTheme()

    def build(self, payload: dict[str, Any]) -> dict[str, Any]:
        conversation = _list(payload.get("conversation") or payload.get("transcript"))
        threshold = _metric_threshold(payload)
        metrics = [_normalize_metric(metric, threshold) for metric in _list(payload.get("metrics")) if isinstance(metric, dict)]
        components = _collect_components(conversation)
        recordings = _list(payload.get("recording_metadata"))
        artifacts = payload.get("artifacts") if isinstance(payload.get("artifacts"), dict) else {}
        score = _float(payload.get("overall_score"))
        passed = payload.get("passed")
        status = _status(payload)
        risks = _risks(payload, metrics, conversation)
        findings = _findings(payload, metrics, components)
        failures = _failure_analysis(payload.get("failure_analysis"))
        failed_message_ids = _failed_message_ids(failures)
        executive = _client_executive_summary(payload, score, status, metrics, failures)
        detailed_failures = _detailed_failure_analysis(failures)
        backend_validation = payload.get("backend_validation") or {}
        rag = payload.get("rag_evaluation") or payload.get("rag")

        generated_at = _now()
        evaluation_id = _value(payload.get("evaluation_id"))
        return {
            "schema": "evalforge.enterprise_report.v1",
            "generated_at": generated_at,
            "brand": self.theme.brand,
            "title": self.theme.title,
            "report_metadata": {
                "client_report_version": CLIENT_REPORT_VERSION,
                "generated_by": self.theme.brand,
                "evalforge_version": EVALFORGE_VERSION,
                "evaluation_id": evaluation_id,
                "generated_at": generated_at,
            },
            "header": {
                "evaluation_id": evaluation_id,
                "agent_name": _value(payload.get("agent_name") or payload.get("connection_name") or payload.get("agent_id") or payload.get("connection_id")),
                "agent_type": _value(payload.get("agent_type")),
                "adapter": _value(payload.get("adapter") or payload.get("provider")),
                "provider": _value(payload.get("provider") or payload.get("adapter")),
                "scenario": _value(payload.get("scenario_id")),
                "date_time": _value(payload.get("created_at") or payload.get("ended_at") or payload.get("started_at")),
                "duration": _value(payload.get("duration_seconds") or payload.get("duration")),
                "status": status,
                "overall_score": score,
                "passed": passed,
            },
            "executive_summary": executive,
            "evaluation_summary": _evaluation_summary(metrics),
            "scorecard": _scorecard(score, status, metrics, threshold),
            "capability_matrix": _capability_matrix(payload, components),
            "conversation_timeline": _conversation_timeline(conversation, failed_message_ids),
            "workflow_timeline": _workflow_timeline(conversation),
            "workflow_validation": _workflow_validation(conversation, metrics),
            "ai_findings": findings,
            "overall_findings": _overall_findings(metrics, failures),
            "risks": risks,
            "overall_risk": executive["overall_risk"],
            "recommendations": _recommendations(payload, metrics, risks),
            "priority_recommendations": _priority_recommendations(failures, metrics, risks),
            "charts": _charts(score, metrics),
            "artifacts": {
                "buttons": _by_type(components, {"button", "quick_reply", "option"}),
                "tables": _by_type(components, {"table"}),
                "forms": _by_type(components, {"form"}),
                "cards": _by_type(components, {"card"}),
                "attachments": _attachments(conversation, components),
                "uploaded_files": _list(artifacts.get("uploaded_files") or payload.get("uploaded_files")),
                "uploaded_images": _list(artifacts.get("uploaded_images") or payload.get("uploaded_images")),
                "uploaded_videos": _list(artifacts.get("uploaded_videos") or payload.get("uploaded_videos")),
                "otp_interactions": _otp(conversation),
                "interaction_summary": _interaction_summary(conversation),
                "recordings": recordings,
            },
            "rag": rag,
            "rag_validation": _rag_validation(rag, metrics, conversation),
            "backend_validation": backend_validation,
            "backend_validation_report": _backend_validation_report(backend_validation),
            "technical_errors": _list(payload.get("technical_errors")),
            "failure_analysis": failures,
            "detailed_failure_analysis": detailed_failures,
            "evidence": payload.get("evidence") or payload.get("conversation_evidence") or [],
            "errors": payload.get("error") or payload.get("errors"),
            "backend_logs": payload.get("backend_logs") or payload.get("logs") or [],
            "appendix": _appendix(payload, failures),
            "raw_report": payload,
        }

    def build_executive(self, payload: dict[str, Any]) -> dict[str, Any]:
        report = self.build(payload)
        report = dict(report)
        report["schema"] = "evalforge.executive_report.v1"
        report["title"] = "Executive AI Agent Assessment Report"
        report["report_type"] = "executive"
        report["report_metadata"] = {
            **report.get("report_metadata", {}),
            "report_type": "Executive Report",
        }
        report["appendix"] = _executive_appendix(report.get("appendix"))
        report.pop("raw_report", None)
        return report

    def build_technical(self, payload: dict[str, Any]) -> dict[str, Any]:
        report = self.build(payload)
        report["schema"] = "evalforge.technical_report.v1"
        report["title"] = "Technical AI Agent Assessment Report"
        report["report_type"] = "technical"
        report["report_metadata"] = {
            **report.get("report_metadata", {}),
            "report_type": "Technical Report",
        }
        return report

    def html(self, payload: dict[str, Any]) -> str:
        return self.technical_html(payload)

    def executive_html(self, payload: dict[str, Any]) -> str:
        return _html_executive_document(self.build_executive(payload), self.theme)

    def technical_html(self, payload: dict[str, Any]) -> str:
        return _html_document(self.build_technical(payload), self.theme)

    def pdf(self, payload: dict[str, Any]) -> bytes:
        return self.technical_pdf(payload)

    def executive_pdf(self, payload: dict[str, Any]) -> bytes:
        return _reportlab_executive_pdf(self.build_executive(payload), self.theme)

    def technical_pdf(self, payload: dict[str, Any]) -> bytes:
        return _reportlab_pdf(self.build_technical(payload), self.theme)

    def docx(self, payload: dict[str, Any]) -> bytes:
        return self.technical_docx(payload)

    def executive_docx(self, payload: dict[str, Any]) -> bytes:
        return _python_docx_executive(self.build_executive(payload), self.theme)

    def technical_docx(self, payload: dict[str, Any]) -> bytes:
        return _python_docx(self.build_technical(payload), self.theme)


def build_enterprise_report(payload: dict[str, Any]) -> dict[str, Any]:
    return EvaluationReportEngine().build(payload)


def build_executive_report(payload: dict[str, Any]) -> dict[str, Any]:
    return EvaluationReportEngine().build_executive(payload)


def build_technical_report(payload: dict[str, Any]) -> dict[str, Any]:
    return EvaluationReportEngine().build_technical(payload)


def build_enterprise_report_html(payload: dict[str, Any]) -> str:
    return EvaluationReportEngine().html(payload)


def build_executive_report_html(payload: dict[str, Any]) -> str:
    return EvaluationReportEngine().executive_html(payload)


def build_technical_report_html(payload: dict[str, Any]) -> str:
    return EvaluationReportEngine().technical_html(payload)


def build_enterprise_report_pdf(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().pdf(payload)


def build_executive_report_pdf(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().executive_pdf(payload)


def build_technical_report_pdf(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().technical_pdf(payload)


def build_enterprise_report_docx(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().docx(payload)


def build_executive_report_docx(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().executive_docx(payload)


def build_technical_report_docx(payload: dict[str, Any]) -> bytes:
    return EvaluationReportEngine().technical_docx(payload)


def _executive_summary(payload: dict[str, Any], score: float | None, status: str, metrics: list[dict[str, Any]], risks: list[dict[str, Any]]) -> dict[str, Any]:
    failed = [m for m in metrics if m.get("passed") is False]
    score_text = _percent(score)
    
    categories = {}
    for m in metrics:
        cat = m.get("category") or "Conversation"
        categories.setdefault(cat, []).append(_float(m.get("score")) or 0.0)
    
    cat_scores = {k: sum(v)/len(v) for k, v in categories.items()}
    
    if payload.get("success") is False:
        narrative = "The evaluation did not complete successfully."
        recommendation = "🔴 High Risk"
    elif status == "PASS":
        narrative = f"The AI agent met the configured assessment threshold with an overall score of {score_text}."
        recommendation = "🟢 Excellent" if (score or 0) > 0.9 else "🟡 Good"
    else:
        narrative = f"The AI agent did not meet the configured assessment threshold. Overall score: {score_text}."
        recommendation = "🟡 Needs Improvement" if (score or 0) > 0.5 else "🔴 High Risk"

    key_outcomes = [
        f"Status: {status}",
        f"Overall Score: {score_text}",
    ]
    if "Conversation" in cat_scores: key_outcomes.append(f"Conversation Score: {_percent(cat_scores['Conversation'])}")
    if "Workflow" in cat_scores: key_outcomes.append(f"Workflow Score: {_percent(cat_scores['Workflow'])}")
    if "rag" in cat_scores: key_outcomes.append(f"RAG Score: {_percent(cat_scores['rag'])}")
    if "Voice" in cat_scores: key_outcomes.append(f"Voice Score: {_percent(cat_scores['Voice'])}")

    return {
        "narrative": narrative,
        "recommendation": recommendation,
        "key_outcomes": key_outcomes,
        "attention_items": [
            _value(m.get("name")) for m in failed[:5]
        ] or ["No failed metric items captured."],
    }


def _client_executive_summary(
    payload: dict[str, Any],
    score: float | None,
    status: str,
    metrics: list[dict[str, Any]],
    failures: list[dict[str, Any]],
) -> dict[str, Any]:
    passed_metrics = [m for m in metrics if m.get("passed") is True]
    failed_metrics = [m for m in metrics if m.get("passed") is False]
    critical_failures = [
        f for f in failures
        if str(f.get("severity") or "").lower() in {"critical", "high"}
    ]
    overall_risk = _overall_risk(score, status, failed_metrics, critical_failures)
    score_text = _percent(score)
    narrative = (
        f"This assessment evaluated {_value(payload.get('agent_name') or payload.get('agent_id') or payload.get('connection_id'))} "
        f"against scenario {_value(payload.get('scenario_id'))}. The evaluation finished with status {status}, "
        f"an overall score of {score_text}, and {len(failed_metrics)} failed metric(s). "
        f"Overall production risk is {overall_risk.lower()}."
    )
    if not failed_metrics and status == "PASS":
        narrative = (
            f"The evaluated AI agent met the assessment criteria with an overall score of {score_text}. "
            "No failed metrics were identified in the completed evaluation."
        )
    return {
        "project": _value(payload.get("project") or "EvalForge AI Agent Evaluation"),
        "agent_name": _value(payload.get("agent_name") or payload.get("connection_name") or payload.get("agent_id") or payload.get("connection_id")),
        "scenario": _value(payload.get("scenario_id")),
        "evaluation_date": _value(payload.get("created_at") or payload.get("ended_at") or payload.get("started_at")),
        "overall_score": score,
        "overall_score_display": score_text,
        "status": status,
        "pass_fail": "PASS" if status == "PASS" else "FAIL",
        "overall_risk": overall_risk,
        "total_metrics": len(metrics),
        "passed": len(passed_metrics),
        "failed": len(failed_metrics),
        "critical_failures": len(critical_failures),
        "evaluation_duration": _value(payload.get("duration_seconds") or payload.get("duration")),
        "narrative": narrative,
        "recommendation": _executive_recommendation(overall_risk, status),
        "key_outcomes": [
            f"Overall Score: {score_text}",
            f"Pass / Fail: {'PASS' if status == 'PASS' else 'FAIL'}",
            f"Overall Risk: {overall_risk}",
            f"Failed Metrics: {len(failed_metrics)}",
            f"Critical Failures: {len(critical_failures)}",
        ],
        "attention_items": [_value(f.get("metric_name")) for f in failures[:5]] or ["No failed metric items captured."],
    }


def _executive_recommendation(risk: str, status: str) -> str:
    if status == "PASS" and risk == "Low":
        return "Ready for controlled rollout with routine monitoring."
    if risk == "High":
        return "Do not proceed to production until priority failures are remediated and the evaluation is rerun."
    return "Proceed only after remediating failed metrics and validating the updated agent behavior."


def _overall_risk(score: float | None, status: str, failed_metrics: list[dict[str, Any]], critical_failures: list[dict[str, Any]]) -> str:
    if status in {"ERROR", "FAIL"} and (score is None or score < 0.5):
        return "High"
    if critical_failures or len(failed_metrics) >= 3:
        return "High"
    if failed_metrics or (score is not None and score < 0.75):
        return "Medium"
    return "Low"


def _evaluation_summary(metrics: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "metric": _value(metric.get("name")),
            "category": _value(metric.get("category") or _metric_category_from_name(metric.get("name"))),
            "score": _float(metric.get("score")),
            "score_display": _percent(_float(metric.get("score"))),
            "threshold": _float(metric.get("threshold")),
            "threshold_display": _percent(_float(metric.get("threshold"))),
            "status": _value(metric.get("status") or ("PASS" if metric.get("passed") else "FAIL")),
            "severity": _metric_severity(metric),
        }
        for metric in metrics
    ]


def _scorecard(score: float | None, status: str, metrics: list[dict[str, Any]], threshold: float | None) -> dict[str, Any]:
    passed = [m for m in metrics if m.get("passed") is True]
    failed = [m for m in metrics if m.get("passed") is False]
    return {
        "overall": {"label": "Overall Score", "score": score, "display": _percent(score), "status": status},
        "pass_threshold": threshold,
        "pass_threshold_display": _percent(threshold),
        "metric_count": len(metrics),
        "passed_metrics": len(passed),
        "failed_metrics": len(failed),
        "metrics": [
            {
                "name": _value(metric.get("name")),
                "score": _float(metric.get("score")),
                "display": _percent(_float(metric.get("score"))),
                "passed": metric.get("passed"),
                "status": _value(metric.get("status")),
                "threshold": _float(metric.get("threshold")),
                "threshold_display": _percent(_float(metric.get("threshold"))),
                "category": _value(metric.get("category")),
                "reason": _value(metric.get("reason")),
            }
            for metric in metrics
        ],
    }


def _capability_matrix(payload: dict[str, Any], components: list[dict[str, Any]]) -> list[dict[str, Any]]:
    def score_to_grade(s: float | None) -> str:
        if s is None: return "Not Evaluated"
        if s >= 0.9: return "Excellent"
        if s >= 0.7: return "Good"
        if s >= 0.5: return "Fair"
        return "Poor"
        
    metrics = payload.get("metrics") or []
    categories = {}
    for m in metrics:
        if isinstance(m, dict):
            cat = m.get("category") or "Conversation"
            if cat == "base":
                cat = "Conversation"
            score = _float(m.get("score"))
            if score is not None:
                categories.setdefault(cat, []).append(score)
            
    cat_scores = {k: sum(v)/len(v) for k, v in categories.items()}
    
    return [
        {"capability": "Conversation", "status": score_to_grade(cat_scores.get("Conversation")), "evidence": f"{len(_list(payload.get('conversation')))} turns"},
        {"capability": "Workflow", "status": score_to_grade(cat_scores.get("Workflow")), "evidence": "Score derived from Workflow Metrics"},
        {"capability": "Knowledge (RAG)", "status": score_to_grade(cat_scores.get("rag")), "evidence": "Score derived from RAG Metrics"},
        {"capability": "Voice", "status": score_to_grade(cat_scores.get("Voice")), "evidence": "Score derived from Voice Metrics"},
        {"capability": "Backend Integration", "status": score_to_grade(cat_scores.get("Backend")), "evidence": "Score derived from Backend Validation"},
    ]



def _workflow_timeline(conversation: list[dict[str, Any]]) -> list[dict[str, Any]]:
    timeline = []
    for index, turn in enumerate(conversation, 1):
        metadata = turn.get("metadata") or {}
        evidence = metadata.get("interaction_evidence")
        if evidence:
            action = _value(evidence.get("action") or evidence.get("interaction_type"))
            success = evidence.get("success", False)
            duration = evidence.get("duration_ms", 0) // 1000 if evidence.get("duration_ms") else 0
            timeline.append({
                "turn": index,
                "action": action,
                "status": "Completed" if success else "Failed",
                "duration": f"{duration} sec" if duration else "N/A",
                "interaction": evidence.get("interaction_type", "interaction"),
                "source": evidence.get("source", "N/A"),
                "timestamp": evidence.get("timestamp", "N/A")
            })
    return timeline


def _conversation_timeline(conversation: list[dict[str, Any]], failed_message_ids: set[str] | None = None) -> list[dict[str, Any]]:
    timeline: list[dict[str, Any]] = []
    failed_message_ids = failed_message_ids or set()
    for index, turn in enumerate(conversation, 1):
        components = _turn_components(turn)
        attachments = _turn_attachments(turn)
        evidence = _interaction_evidence(turn)
        message_id = _message_id(turn)
        timeline.append({
            "turn": index,
            "message_id": message_id,
            "speaker": _value(turn.get("speaker")),
            "message": _value(turn.get("message")),
            "components": components,
            "component_count": len(components),
            "attachments": attachments,
            "attachment_count": len(attachments),
            "interaction_evidence": evidence,
            "failed": bool(message_id and message_id in failed_message_ids),
        })
    return timeline


def _failed_message_ids(failures: list[dict[str, Any]]) -> set[str]:
    ids: set[str] = set()
    for failure in failures:
        message_id = failure.get("conversation_message_id")
        if message_id:
            ids.add(str(message_id))
        refs = failure.get("technical_details") if isinstance(failure.get("technical_details"), dict) else {}
        message_id = (((refs.get("conversation_message_refs") or {}).get("failure_turn") or {}).get("message_id"))
        if message_id:
            ids.add(str(message_id))
    return ids


def _message_id(turn: Any) -> str:
    if not isinstance(turn, dict):
        return ""
    metadata = turn.get("metadata") if isinstance(turn.get("metadata"), dict) else {}
    return _value(turn.get("conversation_message_id") or turn.get("message_id") or turn.get("id") or metadata.get("conversation_message_id") or metadata.get("message_id") or metadata.get("id"))


def _workflow_validation(conversation: list[dict[str, Any]], metrics: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in _workflow_timeline(conversation):
        rows.append({
            "step": item.get("action"),
            "expected": "Interaction should complete successfully.",
            "observed": item.get("status"),
            "status": item.get("status"),
        })
    for metric in metrics:
        category = str(metric.get("category") or "").lower()
        name = str(metric.get("name") or "").lower()
        if "workflow" in category or "workflow" in name:
            rows.append({
                "step": metric.get("name"),
                "expected": f"Score >= {_percent(metric.get('threshold'))}",
                "observed": _percent(metric.get("score")),
                "status": "PASS" if metric.get("passed") else "FAIL",
            })
    return rows


def _rag_validation(rag: Any, metrics: list[dict[str, Any]], conversation: list[dict[str, Any]]) -> dict[str, Any]:
    rag_metrics = [m for m in metrics if str(m.get("category") or "").lower() == "rag" or "rag" in str(m.get("name") or "").lower()]
    questions = []
    last_user = ""
    for turn in conversation:
        if not isinstance(turn, dict):
            continue
        speaker = str(turn.get("speaker") or "").lower()
        if speaker == "user":
            last_user = _value(turn.get("message"))
        elif speaker == "agent":
            rag_ev = turn.get("rag") or (turn.get("metadata") or {}).get("rag_evidence") if isinstance(turn.get("metadata"), dict) else None
            if rag_ev:
                questions.append({
                    "question": last_user,
                    "answer": _value(turn.get("message")),
                    "retrieved_context": _component_value(rag_ev),
                    "score": "N/A",
                    "evidence": _component_value(rag_ev),
                })
    return {
        "summary": rag or {},
        "metrics": rag_metrics,
        "questions": questions,
    }


def _findings(payload: dict[str, Any], metrics: list[dict[str, Any]], components: list[dict[str, Any]]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    for metric in metrics:
        severity = "positive" if metric.get("passed") else "negative"
        findings.append({
            "title": _value(metric.get("name")),
            "severity": severity,
            "summary": _value(metric.get("reason")),
            "score": _percent(_float(metric.get("score"))),
        })
    if components:
        findings.append({
            "title": "Rich Response Evidence Captured",
            "severity": "informational",
            "summary": f"{len(components)} structured UI artifact(s) were preserved in the report.",
            "score": "N/A",
        })
    if payload.get("error"):
        findings.append({
            "title": "Execution Error",
            "severity": "critical",
            "summary": _value(payload.get("error")),
            "score": "N/A",
        })
    return findings or [{"title": "No Findings Captured", "severity": "informational", "summary": "N/A", "score": "N/A"}]


def _risks(payload: dict[str, Any], metrics: list[dict[str, Any]], conversation: list[dict[str, Any]]) -> list[dict[str, Any]]:
    risks: list[dict[str, Any]] = []
    if payload.get("success") is False:
        risks.append({"level": "High", "risk": "Evaluation did not complete successfully.", "impact": "Assessment may be incomplete."})
    for metric in metrics:
        if metric.get("passed") is False:
            risks.append({"level": "Medium", "risk": f"Metric failed: {_value(metric.get('name'))}", "impact": _value(metric.get("reason"))})
    if not conversation:
        risks.append({"level": "High", "risk": "No conversation transcript captured.", "impact": "Report evidence is limited."})
    return risks


def _recommendations(payload: dict[str, Any], metrics: list[dict[str, Any]], risks: list[dict[str, Any]]) -> list[str]:
    recommendations = [
        "Review failed metrics and replay the saved conversation transcript before production approval.",
        "Validate edge-case scenarios against the same connection before expanding rollout.",
    ]
    if payload.get("rag_evaluation"):
        recommendations.append("Review RAG evidence and source grounding for answer faithfulness.")
    if risks:
        recommendations.append("Address high and medium risks, then rerun the evaluation to confirm improvement.")
    if not metrics:
        recommendations.append("Enable or verify metric generation so future reports include quantified assessment evidence.")
    return recommendations


def _priority_recommendations(failures: list[dict[str, Any]], metrics: list[dict[str, Any]], risks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    recommendations: list[dict[str, Any]] = []
    for failure in failures[:5]:
        recommendations.append({
            "priority": len(recommendations) + 1,
            "title": f"Fix {_value(failure.get('metric_name'))}",
            "recommendation": _value(failure.get("recommendation") or "Review the failed evidence and adjust the agent behavior."),
            "severity": _value(failure.get("severity") or "Medium"),
        })
    if not recommendations and risks:
        for risk in risks[:5]:
            recommendations.append({
                "priority": len(recommendations) + 1,
                "title": _value(risk.get("risk")),
                "recommendation": _value(risk.get("impact")),
                "severity": _value(risk.get("level") or "Medium"),
            })
    if not recommendations:
        recommendations.append({
            "priority": 1,
            "title": "Continue monitoring",
            "recommendation": "Maintain evaluation coverage and rerun this scenario after material agent changes.",
            "severity": "Low",
        })
    return recommendations[:5]


def _overall_findings(metrics: list[dict[str, Any]], failures: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = {}
    for metric in metrics:
        category = _value(metric.get("category") or _metric_category_from_name(metric.get("name")))
        groups.setdefault(category, {"category": category, "issues": 0, "metrics": 0})
        groups[category]["metrics"] += 1
        if metric.get("passed") is False:
            groups[category]["issues"] += 1
    for failure in failures:
        category = _value(failure.get("category") or "Failure Analysis")
        groups.setdefault(category, {"category": category, "issues": 0, "metrics": 0})
        if groups[category]["issues"] == 0:
            groups[category]["issues"] += 1
    return list(groups.values())


def _detailed_failure_analysis(failures: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "metric_name": _value(failure.get("metric_name")),
            "severity": _value(failure.get("severity")),
            "category": _value(failure.get("category")),
            "source": _value(failure.get("source")),
            "conversation_turn": failure.get("conversation_turn"),
            "business_summary": _value(failure.get("failure_summary") or failure.get("reason")),
            "expected_behaviour": _value(failure.get("expected")),
            "actual_behaviour": _value(failure.get("actual")),
            "conversation_evidence": _failure_conversation_rows(failure),
            "root_cause": _value(failure.get("root_cause")),
            "backend_evidence": _backend_failure_evidence(failure),
            "supporting_evidence": failure.get("evidence"),
            "business_impact": _value(failure.get("business_impact")),
            "recommendation": _value(failure.get("recommendation")),
            "developer_notes": {
                "metric_score": (failure.get("technical_details") or {}).get("score") if isinstance(failure.get("technical_details"), dict) else None,
                "threshold": (failure.get("technical_details") or {}).get("threshold") if isinstance(failure.get("technical_details"), dict) else None,
                "message_ids": (failure.get("technical_details") or {}).get("conversation_message_refs") if isinstance(failure.get("technical_details"), dict) else None,
                "raw_evidence": failure.get("evidence"),
                "technical_details": failure.get("technical_details"),
            },
        }
        for failure in failures
    ]


def _backend_failure_evidence(failure: dict[str, Any]) -> list[Any]:
    evidence = []
    for item in _list(failure.get("evidence")):
        if isinstance(item, dict) and (
            "backend" in str(item.get("type") or item.get("source") or "").lower()
            or item.get("endpoint")
            or item.get("query")
            or item.get("status_code") is not None
        ):
            evidence.append(item)
    return evidence


def _backend_validation_report(validation: Any) -> list[dict[str, Any]]:
    validation = validation if isinstance(validation, dict) else {}
    checks = validation.get("checks") if isinstance(validation.get("checks"), list) else []
    rows: list[dict[str, Any]] = []
    for check in checks:
        if not isinstance(check, dict):
            continue
        evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
        rows.append({
            "backend": _backend_connection_label(evidence.get("backend_connection")),
            "validation_rule": _value(check.get("name")),
            "expected": check.get("expected"),
            "actual": check.get("actual"),
            "status": "PASS" if check.get("passed") else "FAIL",
            "evidence": evidence,
        })
    return rows


def _appendix(payload: dict[str, Any], failures: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "technical_details": payload.get("technical_details") or {},
        "api_logs": payload.get("backend_logs") or payload.get("logs") or [],
        "conversation_ids": [
            {
                "turn": turn.get("turn_index") or index,
                "message_id": _message_id(turn),
                "speaker": turn.get("speaker"),
            }
            for index, turn in enumerate(_list(payload.get("conversation") or payload.get("transcript")))
            if isinstance(turn, dict)
        ],
        "evidence_ids": [failure.get("failure_id") for failure in failures if failure.get("failure_id")],
        "attachments": payload.get("attachments") or [],
        "screenshots": payload.get("screenshots") or [],
        "tool_calls": payload.get("tool_calls") or [],
        "backend_requests": payload.get("backend_requests") or [],
        "stack_traces": payload.get("stack_traces") or [],
        "raw_json": _sanitize_report_json(payload),
    }


def _executive_appendix(appendix: Any) -> dict[str, Any]:
    appendix = appendix if isinstance(appendix, dict) else {}
    return {
        "conversation_messages": len(_list(appendix.get("conversation_ids"))),
        "evidence_items": len(_list(appendix.get("evidence_ids"))),
        "attachments": len(_list(appendix.get("attachments"))),
        "screenshots": len(_list(appendix.get("screenshots"))),
        "backend_requests": len(_list(appendix.get("backend_requests"))),
        "technical_appendix_available": True,
    }


def _charts(score: float | None, metrics: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "overall_score": {"label": "Overall Score", "value": score or 0},
        "metric_scores": [
            {"label": _value(metric.get("name")), "value": _float(metric.get("score")) or 0, "passed": bool(metric.get("passed"))}
            for metric in metrics
        ],
    }


def _html_executive_document(report: dict[str, Any], theme: ReportTheme) -> str:
    header = report["header"]
    status_class = "pass" if header["passed"] else ("error" if report["errors"] else "fail")
    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>{e(report["title"])} - {e(header["evaluation_id"])}</title>
  <style>{_html_css(theme)}</style>
</head>
<body>
  <main class="report-shell executive-report">
    <section class="cover">
      <div class="brand">{e(theme.brand)}</div>
      <h1>{e(report["title"])}</h1>
      <p>{e(report["executive_summary"]["narrative"])}</p>
      <span class="status {status_class}">{e(header["status"])}</span>
      <div class="meta-grid">
        {_meta("Client", report["executive_summary"].get("project"))}
        {_meta("Agent", header["agent_name"])}
        {_meta("Scenario", header["scenario"])}
        {_meta("Evaluation Date", header["date_time"])}
        {_meta("Overall Score", _percent(header["overall_score"]))}
        {_meta("Overall Risk", report["executive_summary"]["overall_risk"])}
        {_meta("Evaluation ID", header["evaluation_id"])}
        {_meta("Version", report["report_metadata"].get("evalforge_version"))}
        {_meta("Generated By", report["report_metadata"].get("generated_by"))}
      </div>
      {_html_report_metadata(report)}
    </section>
    {_html_executive_summary(report)}
    {_html_executive_kpi_dashboard(report)}
    {_html_capability_matrix(report)}
    {_html_executive_failure_analysis(report)}
    {_html_backend_validation_business(report)}
    {_html_workflow_validation(report)}
    {_html_rag_validation(report)}
    {_html_overall_findings(report)}
    {_html_priority_recommendations(report)}
    {_html_final_recommendation(report)}
  </main>
</body>
</html>"""


def _html_document(report: dict[str, Any], theme: ReportTheme) -> str:
    header = report["header"]
    status_class = "pass" if header["passed"] else ("error" if report["errors"] else "fail")
    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>{e(report.get("title") or theme.title)} - {e(header["evaluation_id"])}</title>
  <style>{_html_css(theme)}</style>
</head>
<body>
  <main class="report-shell">
    <section class="cover">
      <div class="brand">{e(theme.brand)}</div>
      <h1>{e(report.get("title") or theme.title)}</h1>
      <p>{e(report["executive_summary"]["narrative"])}</p>
      <p><strong>Management Recommendation:</strong> {e(report["executive_summary"].get("recommendation"))}</p>
      <span class="status {status_class}">{e(header["status"])}</span>
      <div class="meta-grid">
        {_meta("Project", report["executive_summary"]["project"])}
        {_meta("Agent", header["agent_name"])}
        {_meta("Scenario", header["scenario"])}
        {_meta("Adapter", header["adapter"])}
        {_meta("Date & Time", header["date_time"])}
        {_meta("Duration", header["duration"])}
        {_meta("Score", _percent(header["overall_score"]))}
        {_meta("Overall Risk", report["executive_summary"]["overall_risk"])}
        {_meta("Total Metrics", report["executive_summary"]["total_metrics"])}
        {_meta("Critical Failures", report["executive_summary"]["critical_failures"])}
      </div>
      {_html_report_metadata(report)}
    </section>
    {_html_executive_summary(report)}
    {_html_evaluation_summary(report)}
    {_html_scorecard(report)}
    {_html_capability_matrix(report)}
    {_html_detailed_failure_analysis(report)}
    {_html_timeline(report)}
    {_html_backend_validation(report)}
    {_html_workflow_validation(report)}
    {_html_rag_validation(report)}
    {_html_overall_findings(report)}
    {_html_priority_recommendations(report)}
    {_html_artifacts(report)}
    {_html_json_section("Appendix", report["appendix"])}
  </main>
</body>
</html>"""


def _html_css(theme: ReportTheme) -> str:
    return f"""
body {{ margin: 0; background: #eef2f7; color: {theme.ink}; font-family: Inter, Arial, sans-serif; }}
.report-shell {{ max-width: 1180px; margin: 0 auto; padding: 32px; }}
.cover, .section {{ background: white; border: 1px solid #dbe3ef; border-radius: 8px; padding: 24px; margin-bottom: 18px; box-shadow: 0 10px 30px rgba(15,23,42,.06); }}
.brand {{ color: {theme.primary}; font-weight: 900; letter-spacing: .12em; text-transform: uppercase; font-size: 12px; }}
h1 {{ margin: 8px 0; font-size: 34px; }} h2 {{ margin: 0 0 14px; font-size: 20px; }}
p {{ line-height: 1.65; }} .muted {{ color: {theme.muted}; }}
.status {{ display: inline-block; padding: 7px 12px; border-radius: 999px; font-weight: 800; font-size: 12px; }}
.pass {{ background: #dcfce7; color: {theme.success}; }} .fail {{ background: #fee2e2; color: {theme.danger}; }} .error {{ background: #ffedd5; color: {theme.warning}; }}
.meta-grid, .score-grid, .artifact-grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 12px; margin-top: 18px; }}
.report-metadata {{ grid-template-columns: repeat(5, 1fr); }}
.meta, .score-card, .artifact-card {{ border: 1px solid #e2e8f0; background: #f8fafc; border-radius: 8px; padding: 14px; }}
.label {{ color: {theme.muted}; font-size: 11px; text-transform: uppercase; letter-spacing: .08em; }}
.value {{ margin-top: 6px; font-weight: 800; overflow-wrap: anywhere; }}
.bar {{ height: 9px; background: #e2e8f0; border-radius: 999px; overflow: hidden; margin-top: 10px; }}
.bar span {{ display: block; height: 100%; background: {theme.primary}; }}
table {{ width: 100%; border-collapse: collapse; font-size: 12px; }} th, td {{ border: 1px solid #e2e8f0; padding: 8px; text-align: left; vertical-align: top; }} th {{ background: #f1f5f9; }}
.timeline {{ border-left: 3px solid #dbeafe; padding-left: 16px; }} .turn {{ margin: 12px 0; padding: 12px; background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; }}
.turn.failed-turn {{ border-color: #fecaca; background: #fff7f7; box-shadow: inset 4px 0 0 {theme.danger}; }}
.speaker {{ font-size: 11px; color: {theme.muted}; font-weight: 900; text-transform: uppercase; margin-bottom: 5px; }}
.components {{ margin-top: 10px; display: grid; gap: 8px; }}
.button-pill {{ display: inline-block; border: 1px solid #bfdbfe; background: #eff6ff; color: #1d4ed8; border-radius: 999px; padding: 6px 10px; margin: 3px 5px 3px 0; font-weight: 700; font-size: 12px; }}
.component-panel, .attachment-card {{ border: 1px solid #dbe3ef; background: white; border-radius: 8px; padding: 10px; }}
.component-title {{ font-weight: 800; margin-bottom: 6px; color: {theme.ink}; }}
blockquote {{ margin: 0 0 10px; padding: 10px 12px; border-left: 3px solid #bfdbfe; background: white; border-radius: 6px; }}
.rich-image, .rich-video {{ max-width: 100%; border-radius: 8px; border: 1px solid #e2e8f0; margin-top: 6px; }}
.artifact-json {{ max-height: 180px; }}
.finding, .risk, .recommendation {{ border: 1px solid #e2e8f0; border-radius: 8px; padding: 12px; margin: 10px 0; }}
.failure-card {{ border: 1px solid #fecaca; background: #fff7f7; border-radius: 8px; padding: 16px; margin: 12px 0; }}
.failure-header {{ display: flex; justify-content: space-between; gap: 12px; align-items: flex-start; flex-wrap: wrap; }}
.failure-title {{ font-weight: 900; font-size: 16px; }}
.failure-badges {{ display: flex; gap: 8px; flex-wrap: wrap; }}
.failure-badge {{ display: inline-block; border-radius: 999px; padding: 5px 9px; font-size: 11px; font-weight: 900; text-transform: uppercase; }}
.failure-badge.status {{ background: #fee2e2; color: {theme.danger}; }}
.failure-badge.severity {{ background: #ffedd5; color: {theme.warning}; }}
.failure-badge.category {{ background: #dbeafe; color: {theme.primary}; }}
.failure-grid {{ display: grid; grid-template-columns: repeat(2, 1fr); gap: 10px; margin-top: 12px; }}
.failure-field {{ border: 1px solid #e2e8f0; background: white; border-radius: 8px; padding: 10px; }}
.failure-field.full {{ grid-column: 1 / -1; }}
.failure-snippet {{ border-left: 3px solid #cbd5e1; background: white; border-radius: 8px; padding: 10px; margin: 8px 0; }}
.failure-snippet.failure {{ border-left-color: {theme.danger}; background: #fef2f2; }}
.failure-marker {{ color: {theme.danger}; font-weight: 900; font-size: 11px; text-transform: uppercase; }}
.failure-details {{ margin-top: 10px; }}
.failure-details summary {{ cursor: pointer; font-weight: 800; color: {theme.ink}; }}
.two-col {{ display: grid; grid-template-columns: repeat(2, 1fr); gap: 16px; }}
.page-section {{ break-before: page; page-break-before: always; }}
pre {{ background: #0f172a; color: #e2e8f0; padding: 14px; border-radius: 8px; overflow: auto; max-height: 360px; }}
@media print {{ .section, .cover {{ box-shadow: none; break-inside: avoid; }} .page-section {{ break-before: page; }} }}
@media(max-width: 780px) {{ .meta-grid, .score-grid, .artifact-grid, .two-col {{ grid-template-columns: 1fr; }} .report-shell {{ padding: 16px; }} }}
"""


def _html_scorecard(report: dict[str, Any]) -> str:
    scorecard = report["scorecard"]
    metric_cards = "".join(
        f'<div class="score-card"><div class="label">{e(metric["name"])}</div><div class="value">{e(metric["display"])} · {e(metric["status"])}</div><div class="bar"><span style="width:{_bar_width(metric["score"])}%"></span></div><p class="muted">Threshold: {e(metric["threshold_display"])} · {e(metric["reason"])}</p></div>'
        for metric in scorecard["metrics"]
    ) or '<div class="score-card"><div class="value">N/A</div></div>'
    return f'<section class="section"><h2>Executive Scorecard</h2><div class="score-grid">{_meta("Overall", scorecard["overall"]["display"])}{_meta("Pass Threshold", scorecard["pass_threshold_display"])}{_meta("Passed Metrics", scorecard["passed_metrics"])}{_meta("Failed Metrics", scorecard["failed_metrics"])}</div><div class="score-grid">{metric_cards}</div></section>'


def _html_executive_kpi_dashboard(report: dict[str, Any]) -> str:
    kpis = _executive_kpis(report)
    cards = "".join(
        f'<div class="score-card"><div class="label">{e(item["label"])}</div><div class="value">{e(item["value"])}</div><p class="muted">{e(item.get("note"))}</p></div>'
        for item in kpis
    )
    return f'<section class="section"><h2>Page 3 · KPI Dashboard</h2><div class="score-grid">{cards}</div></section>'


def _executive_kpis(report: dict[str, Any]) -> list[dict[str, str]]:
    scorecard = report.get("scorecard") or {}
    metrics = scorecard.get("metrics") or []
    by_category: dict[str, list[float]] = {}
    for metric in metrics:
        category = _value(metric.get("category") or _metric_category_from_name(metric.get("name")))
        score = _float(metric.get("score"))
        if score is not None:
            by_category.setdefault(category.lower(), []).append(score)

    def category_score(*names: str) -> str:
        values: list[float] = []
        for name in names:
            values.extend(by_category.get(name.lower(), []))
        return _percent(sum(values) / len(values)) if values else "N/A"

    conversation = report.get("conversation_timeline") or []
    duration = _value((report.get("header") or {}).get("duration"))
    return [
        {"label": "Overall Score", "value": _percent((report.get("header") or {}).get("overall_score")), "note": _value((report.get("header") or {}).get("status"))},
        {"label": "Conversation Score", "value": category_score("Conversation", "Response Validation"), "note": "Dialog quality and response behavior."},
        {"label": "Workflow Score", "value": category_score("Workflow"), "note": "Task completion and process adherence."},
        {"label": "Backend Score", "value": category_score("Backend", "Backend Validation"), "note": "Business system validation."},
        {"label": "RAG Score", "value": category_score("RAG", "rag"), "note": "Grounding and retrieval quality."},
        {"label": "Safety Score", "value": category_score("Safety"), "note": "Policy and risk controls."},
        {"label": "Latency", "value": duration, "note": "Evaluation duration."},
        {"label": "Pass Rate", "value": _percent((scorecard.get("passed_metrics") or 0) / max(scorecard.get("metric_count") or 1, 1)), "note": f"{scorecard.get('passed_metrics') or 0} of {scorecard.get('metric_count') or 0} metrics passed."},
        {"label": "Transcript Size", "value": str(len(conversation)), "note": "Captured conversation turns."},
    ]


def _executive_strengths(report: dict[str, Any]) -> str:
    metrics = ((report.get("scorecard") or {}).get("metrics")) or []
    passed = [metric.get("name") for metric in metrics if metric.get("passed") is True]
    if passed:
        return ", ".join(_value(item) for item in passed[:4])
    return "No passed metric strengths were captured."


def _executive_weaknesses(report: dict[str, Any]) -> str:
    failures = report.get("failure_analysis") or []
    if failures:
        return ", ".join(_value(item.get("metric_name")) for item in failures[:4] if isinstance(item, dict))
    metrics = ((report.get("scorecard") or {}).get("metrics")) or []
    failed = [metric.get("name") for metric in metrics if metric.get("passed") is False]
    if failed:
        return ", ".join(_value(item) for item in failed[:4])
    return "No material weaknesses were captured."


def _final_recommendation_label(report: dict[str, Any]) -> str:
    summary = report.get("executive_summary") or {}
    risk = _value(summary.get("overall_risk"))
    status = _value(summary.get("status"))
    if status == "PASS" and risk == "Low":
        return "Deploy"
    if risk == "High":
        return "Do Not Deploy"
    return "Deploy with Monitoring"


def _executive_findings_narrative(report: dict[str, Any]) -> str:
    summary = report.get("executive_summary") or {}
    failed = summary.get("failed") or 0
    risk = summary.get("overall_risk") or "N/A"
    return (
        f"The evaluation identified {failed} failed metric(s), with an overall risk level of {risk}. "
        f"The recommended decision is {_final_recommendation_label(report).lower()}. "
        "Priority remediation should focus on failed metrics with direct business impact before production expansion."
    )


def _html_executive_failure_analysis(report: dict[str, Any]) -> str:
    failures = [item for item in _list(report.get("detailed_failure_analysis")) if isinstance(item, dict)]
    if not failures:
        return ""
    cards = []
    for failure in failures:
        cards.append(
            f'''<section class="section page-section failure-card">
                <div class="failure-header">
                  <div><div class="label">Failure Analysis</div><div class="failure-title">{e(failure.get("metric_name"))}</div></div>
                  <div class="failure-badges"><span class="failure-badge severity">{e(failure.get("severity"))}</span><span class="failure-badge category">{e(failure.get("category"))}</span></div>
                </div>
                <div class="failure-grid">
                  {_failure_text_block("Business Summary", failure.get("business_summary"))}
                  {_failure_text_block("Expected Behaviour", failure.get("expected_behaviour"))}
                  {_failure_text_block("Actual Behaviour", failure.get("actual_behaviour"))}
                  {_failure_text_block("Business Impact", failure.get("business_impact"))}
                  {_html_failure_conversation_rows(failure.get("conversation_evidence"))}
                  {_failure_text_block("Recommendation", failure.get("recommendation"))}
                </div>
            </section>'''
        )
    return "".join(cards)


def _html_backend_validation_business(report: dict[str, Any]) -> str:
    validation = report.get("backend_validation")
    rows = _backend_validation_rows(validation)
    if len(rows) <= 1:
        return ""
    body = "".join(
        f"<tr><td>{e(row[0])}</td><td>{e(row[1])}</td><td>{e(row[2]) if len(row) > 2 else 'N/A'}</td><td>{e(row[3]) if len(row) > 3 else 'N/A'}</td></tr>"
        for row in rows[1:]
    )
    return f'<section class="section"><h2>Backend Validation</h2><p>Backend checks compare the agent outcome against configured business systems.</p><table><thead><tr><th>Expected API</th><th>Actual API</th><th>Result</th><th>Recommendation</th></tr></thead><tbody>{body}</tbody></table></section>'


def _html_final_recommendation(report: dict[str, Any]) -> str:
    summary = report.get("executive_summary") or {}
    risk = _value(summary.get("overall_risk"))
    status = _value(summary.get("status"))
    if status == "PASS" and risk == "Low":
        decision = "Deploy"
    elif risk == "High":
        decision = "Do Not Deploy"
    else:
        decision = "Deploy with Monitoring"
    return f'<section class="section"><h2>Final Recommendation</h2><div class="score-card"><div class="label">Recommendation</div><div class="value">{e(decision)}</div><p>{e(summary.get("recommendation"))}</p></div></section>'


def _html_executive_summary(report: dict[str, Any]) -> str:
    summary = report["executive_summary"]
    cards = "".join(
        _meta(label, summary.get(key))
        for label, key in (
            ("Project", "project"),
            ("Overall Risk", "overall_risk"),
            ("Passed", "passed"),
            ("Failed", "failed"),
            ("Evaluation Duration", "evaluation_duration"),
            ("Pass / Fail", "pass_fail"),
        )
    )
    outcomes = "".join(f"<li>{e(item)}</li>" for item in _list(summary.get("key_outcomes")))
    attention = "".join(f"<li>{e(item)}</li>" for item in _list(summary.get("attention_items")))
    return f'''<section class="section executive-summary">
        <h2>Section 1 · Executive Summary</h2>
        <p>{e(summary.get("narrative"))}</p>
        <div class="meta-grid">{cards}</div>
        <div class="two-col">
            <div><h3>Key Outcomes</h3><ul>{outcomes}</ul></div>
            <div><h3>Attention Items</h3><ul>{attention}</ul></div>
        </div>
    </section>'''


def _html_report_metadata(report: dict[str, Any]) -> str:
    metadata = report.get("report_metadata") or {}
    return '<div class="meta-grid report-metadata">' + "".join([
        _meta("Client Report Version", metadata.get("client_report_version")),
        _meta("Generated By", metadata.get("generated_by")),
        _meta("EvalForge Version", metadata.get("evalforge_version")),
        _meta("Evaluation ID", metadata.get("evaluation_id")),
        _meta("Generated At", metadata.get("generated_at")),
    ]) + "</div>"


def _html_evaluation_summary(report: dict[str, Any]) -> str:
    rows = "".join(
        f"<tr><td>{e(item['metric'])}</td><td>{e(item['category'])}</td><td>{e(item['score_display'])}</td><td>{e(item['threshold_display'])}</td><td>{e(item['status'])}</td><td>{e(item['severity'])}</td></tr>"
        for item in report.get("evaluation_summary") or []
    ) or '<tr><td colspan="6" class="muted">N/A</td></tr>'
    return f'''<section class="section">
        <h2>Section 2 · Evaluation Summary</h2>
        <table><thead><tr><th>Metric</th><th>Category</th><th>Score</th><th>Threshold</th><th>Status</th><th>Severity</th></tr></thead><tbody>{rows}</tbody></table>
    </section>'''


def _html_capability_matrix(report: dict[str, Any]) -> str:
    rows = "".join(f'<tr><td>{e(row["capability"])}</td><td>{e(row["status"])}</td><td>{e(row["evidence"])}</td></tr>' for row in report["capability_matrix"])
    return f'<section class="section"><h2>Capability Matrix</h2><table><thead><tr><th>Capability</th><th>Status</th><th>Evidence</th></tr></thead><tbody>{rows}</tbody></table></section>'


def _html_findings(report: dict[str, Any]) -> str:
    items = "".join(f'<div class="finding"><strong>{e(item["title"])}</strong><div class="muted">{e(item["severity"])} · {e(item["score"])}</div><p>{e(item["summary"])}</p></div>' for item in report["ai_findings"])
    return f'<section class="section"><h2>AI Findings</h2>{items}</section>'



def _html_workflow(report: dict[str, Any]) -> str:
    timeline = report.get("workflow_timeline")
    if not timeline:
        return ""

    rows = "".join(
        f"<tr><td>Turn {item['turn']}</td><td>{e(item['action'])}</td><td>{e(item['status'])}</td><td>{e(item['duration'])}</td></tr>"
        for item in timeline
    )
    
    evidence_rows = "".join(
        f"<tr><td>{e(item['interaction'])}</td><td>{e(item['status'])}</td><td>{e(item['duration'])}</td><td>Interaction #{item['turn']}</td></tr>"
        for item in timeline
    )

    return f'''<section class="section">
        <h2>Workflow Summary</h2>
        <table class="data-table">
            <thead><tr><th>Turn</th><th>Action</th><th>Status</th><th>Duration</th></tr></thead>
            <tbody>{rows}</tbody>
        </table>
        <h3>Evidence Table</h3>
        <table class="data-table">
            <thead><tr><th>Interaction</th><th>Result</th><th>Duration</th><th>Timestamp</th></tr></thead>
            <tbody>{evidence_rows}</tbody>
        </table>
    </section>'''


def _html_workflow_validation(report: dict[str, Any]) -> str:
    rows = report.get("workflow_validation") or []
    if not rows:
        return ""
    body = "".join(
        f"<tr><td>{e(row.get('step'))}</td><td>{e(row.get('expected'))}</td><td>{e(row.get('observed'))}</td><td>{e(row.get('status'))}</td></tr>"
        for row in rows
    )
    return f'''<section class="section">
        <h2>Section 6 · Workflow Validation</h2>
        <table><thead><tr><th>Step</th><th>Expected</th><th>Observed</th><th>Status</th></tr></thead><tbody>{body}</tbody></table>
    </section>'''


def _html_rag_validation(report: dict[str, Any]) -> str:
    rag = report.get("rag_validation") or {}
    if not rag.get("summary") and not rag.get("metrics") and not rag.get("questions"):
        return ""
    metric_rows = "".join(
        f"<tr><td>{e(metric.get('name'))}</td><td>{e(_percent(metric.get('score')))}</td><td>{e(metric.get('status'))}</td><td>{e(metric.get('reason'))}</td></tr>"
        for metric in _list(rag.get("metrics"))
    )
    question_rows = "".join(
        f"<tr><td>{e(item.get('question'))}</td><td>{e(item.get('retrieved_context'))}</td><td>{e(item.get('answer'))}</td><td>{e(item.get('score'))}</td><td>{e(item.get('evidence'))}</td></tr>"
        for item in _list(rag.get("questions"))
    )
    return f'''<section class="section">
        <h2>Section 7 · RAG Validation</h2>
        <h3>RAG Metrics</h3>
        <table><thead><tr><th>Metric</th><th>Score</th><th>Status</th><th>Evidence</th></tr></thead><tbody>{metric_rows or '<tr><td colspan="4">N/A</td></tr>'}</tbody></table>
        <h3>RAG Evidence</h3>
        <table><thead><tr><th>Question</th><th>Retrieved Context</th><th>Answer</th><th>Score</th><th>Evidence</th></tr></thead><tbody>{question_rows or '<tr><td colspan="5">N/A</td></tr>'}</tbody></table>
    </section>'''


def _html_overall_findings(report: dict[str, Any]) -> str:
    rows = "".join(
        f"<tr><td>{e(item.get('category'))}</td><td>{e(item.get('issues'))}</td><td>{e(item.get('metrics'))}</td></tr>"
        for item in report.get("overall_findings") or []
    ) or '<tr><td colspan="3">N/A</td></tr>'
    return f'''<section class="section">
        <h2>Section 8 · Overall Findings</h2>
        <table><thead><tr><th>Category</th><th>Issues</th><th>Metrics Evaluated</th></tr></thead><tbody>{rows}</tbody></table>
    </section>'''


def _html_priority_recommendations(report: dict[str, Any]) -> str:
    items = "".join(
        f'<div class="recommendation"><strong>Priority {e(item.get("priority"))}: {e(item.get("title"))}</strong><p>{e(item.get("recommendation"))}</p><div class="muted">Severity: {e(item.get("severity"))}</div></div>'
        for item in report.get("priority_recommendations") or []
    )
    return f'<section class="section"><h2>Section 9 · Priority Recommendations</h2>{items or "<p>N/A</p>"}</section>'


def _html_backend_validation(report: dict[str, Any]) -> str:
    validation = report.get("backend_validation") or {}
    if not validation:
        return ""
    checks = validation.get("checks") if isinstance(validation.get("checks"), list) else []
    status = _value(validation.get("status") or "not_configured")
    reason = _value(validation.get("reason") or "N/A")
    timeline = _html_backend_validation_timeline(validation)
    explorer = _html_backend_validation_explorer(validation)
    if not timeline and not explorer:
        rows = "".join(_html_backend_validation_check_row(check) for check in checks if isinstance(check, dict))
        if not rows:
            rows = '<tr><td colspan="5" class="muted">N/A</td></tr>'
        explorer = f'''<table class="data-table">
            <thead><tr><th>Check</th><th>Result</th><th>Reason</th><th>Expected</th><th>Actual</th></tr></thead>
            <tbody>{rows}</tbody>
        </table>'''
    return f'''<section class="section">
        <h2>Section 5 · Backend Validation</h2>
        <p><strong>Status:</strong> {e(status)}</p>
        <p class="muted">{e(reason)}</p>
        {timeline}
        {explorer}
    </section>'''


def _html_backend_validation_timeline(validation: dict[str, Any]) -> str:
    events = _backend_validation_timeline_events(validation)
    if not events:
        return ""
    rows = "".join(
        f"<tr><td>{idx}</td><td>{e(event['title'])}</td><td>{e(event['status'])}</td><td>{e(event['detail'])}</td><td>{e(event.get('meta') or 'N/A')}</td></tr>"
        for idx, event in enumerate(events, start=1)
    )
    return f'''<h3>Backend Validation Timeline</h3>
        <table class="data-table">
            <thead><tr><th>#</th><th>Stage</th><th>Status</th><th>Detail</th><th>Evidence</th></tr></thead>
            <tbody>{rows}</tbody>
        </table>'''


def _html_backend_validation_explorer(validation: dict[str, Any]) -> str:
    checks = validation.get("checks") if isinstance(validation.get("checks"), list) else []
    if not checks:
        return ""
    sections = []
    for check in checks:
        if not isinstance(check, dict):
            continue
        evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
        assertion = evidence.get("assertion") if isinstance(evidence.get("assertion"), dict) else {}
        target = evidence.get("endpoint") or evidence.get("query") or "N/A"
        rows = [
            ["Validator", check.get("name") or "Check"],
            ["Backend Connection Used", _backend_connection_label(evidence.get("backend_connection"))],
            ["Endpoint / Query", f"{evidence.get('method') or ''} {target}".strip()],
            ["Latency", f"{evidence.get('response_time_ms')} ms" if evidence.get("response_time_ms") is not None else "N/A"],
            ["Status Code", evidence.get("status_code") if evidence.get("status_code") is not None else "N/A"],
            ["Response Size", f"{len(str(evidence.get('response_body')))} chars" if evidence.get("response_body") is not None else "N/A"],
            ["Variables Used", _component_value(evidence.get("extracted_variables") or {})],
            ["Assertion", f"{assertion.get('jsonpath') or 'N/A'} {assertion.get('operator') or ''}".strip()],
            ["Expected", _component_value(assertion.get("expected", check.get("expected")))],
            ["Actual", _component_value(assertion.get("actual", check.get("actual")))],
        ]
        body = "".join(f"<tr><td>{e(row[0])}</td><td>{e(row[1])}</td></tr>" for row in rows)
        result = "PASS" if check.get("passed") else "FAIL"
        sections.append(
            f'''<details class="report-section" open>
                <summary>{e(check.get('name') or 'Backend validator')} — {e(result)}</summary>
                <table class="data-table"><tbody>{body}</tbody></table>
            </details>'''
        )
    return "<h3>Backend Explorer</h3>" + "".join(sections)


def _backend_validation_timeline_events(validation: dict[str, Any]) -> list[dict[str, Any]]:
    checks = validation.get("checks") if isinstance(validation.get("checks"), list) else []
    if not validation:
        return []
    events = [
        {
            "title": "Conversation completed",
            "status": "DONE",
            "detail": "Backend validation started after the completed transcript.",
        }
    ]
    variables: dict[str, Any] = {}
    for check in checks:
        if not isinstance(check, dict):
            continue
        evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
        extracted = evidence.get("extracted_variables") if isinstance(evidence.get("extracted_variables"), dict) else {}
        variables.update(extracted)
    if variables:
        events.append({
            "title": "Variable extraction",
            "status": "DONE",
            "detail": ", ".join(f"{key}={value}" for key, value in variables.items()),
        })
    for check in checks:
        if not isinstance(check, dict):
            continue
        evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
        status = "PASS" if check.get("passed") else "FAIL"
        if evidence.get("endpoint") or evidence.get("method"):
            events.append({
                "title": "REST validator execution",
                "status": status,
                "detail": f"{evidence.get('method') or 'GET'} {evidence.get('endpoint') or 'N/A'}",
                "meta": _backend_evidence_meta(evidence),
            })
        elif evidence.get("query"):
            events.append({
                "title": "SQL validator execution",
                "status": status,
                "detail": evidence.get("query") or "N/A",
                "meta": _backend_evidence_meta(evidence),
            })
        elif not evidence.get("assertion"):
            events.append({
                "title": "Backend validator execution",
                "status": status,
                "detail": check.get("name") or "Validation check",
                "meta": check.get("reason") or "",
            })
        assertion = evidence.get("assertion") if isinstance(evidence.get("assertion"), dict) else None
        if assertion:
            events.append({
                "title": "Assertion execution",
                "status": assertion.get("status") or status,
                "detail": f"{assertion.get('jsonpath') or 'value'} {assertion.get('operator') or ''}".strip(),
                "meta": f"Expected: {_component_value(assertion.get('expected'))}; Actual: {_component_value(assertion.get('actual'))}",
            })
    events.append({
        "title": "Final backend result",
        "status": validation.get("status") or "N/A",
        "detail": validation.get("reason") or "N/A",
    })
    return events


def _backend_evidence_meta(evidence: dict[str, Any]) -> str:
    parts = []
    if evidence.get("response_time_ms") is not None:
        parts.append(f"Latency: {evidence.get('response_time_ms')} ms")
    if evidence.get("status_code") is not None:
        parts.append(f"Status: {evidence.get('status_code')}")
    if evidence.get("row_count") is not None:
        parts.append(f"Rows: {evidence.get('row_count')}")
    if evidence.get("response_body") is not None:
        parts.append(f"Response: {len(str(evidence.get('response_body')))} chars")
    return " | ".join(parts)


def _html_backend_validation_check_row(check: dict[str, Any]) -> str:
    evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
    endpoint = ""
    if evidence.get("method") or evidence.get("endpoint"):
        endpoint = f"<br><span class=\"muted\">Endpoint: {e(evidence.get('method') or 'GET')} {e(evidence.get('endpoint') or 'N/A')}</span>"
    latency = ""
    if evidence.get("response_time_ms") is not None:
        latency = f"<br><span class=\"muted\">Latency: {e(evidence.get('response_time_ms'))} ms</span>"
    status = "PASS" if check.get("passed") else "FAIL"
    return (
        f"<tr><td>{e(check.get('name'))}{endpoint}{latency}</td>"
        f"<td>{e(status)}</td><td>{e(check.get('reason'))}</td>"
        f"<td>{e(_component_value(check.get('expected')))}</td>"
        f"<td>{e(_component_value(check.get('actual')))}</td></tr>"
    )


def _html_failure_analysis(report: dict[str, Any]) -> str:
    failures = [item for item in _list(report.get("failure_analysis")) if isinstance(item, dict)]
    if not failures:
        return ""
    cards = "".join(_html_failure_card(failure, index) for index, failure in enumerate(failures, start=1))
    return f'<section class="section"><h2>Failure Analysis</h2>{cards}</section>'


def _html_detailed_failure_analysis(report: dict[str, Any]) -> str:
    failures = [item for item in _list(report.get("detailed_failure_analysis")) if isinstance(item, dict)]
    if not failures:
        return ""
    pages = "".join(_html_detailed_failure_page(failure, index) for index, failure in enumerate(failures, start=1))
    return f'<section class="section detailed-failures"><h2>Section 3 · Detailed Failure Analysis</h2>{pages}</section>'


def _html_detailed_failure_page(failure: dict[str, Any], index: int) -> str:
    backend = failure.get("backend_evidence") or []
    backend_html = _html_failure_evidence(backend) if backend else ""
    supporting_html = _html_failure_evidence(failure.get("supporting_evidence")) if failure.get("supporting_evidence") else ""
    dev_notes = failure.get("developer_notes") if isinstance(failure.get("developer_notes"), dict) else {}
    return f'''<article class="failure-card page-section">
        <div class="failure-header">
            <div>
                <div class="failure-title">❌ {e(failure.get("metric_name") or f"Failure {index}")}</div>
                <p class="muted">Source: {e(failure.get("source"))} · Conversation Turn: {e(failure.get("conversation_turn"))}</p>
            </div>
            <div class="failure-badges">
                <span class="failure-badge severity">{e(failure.get("severity"))}</span>
                <span class="failure-badge category">{e(failure.get("category"))}</span>
            </div>
        </div>
        {_failure_text_block("Business Summary", failure.get("business_summary"))}
        {_failure_text_block("Expected Behaviour", failure.get("expected_behaviour"))}
        {_failure_text_block("Actual Behaviour", failure.get("actual_behaviour"))}
        {_html_failure_conversation_rows(failure.get("conversation_evidence"))}
        {_failure_text_block("Root Cause", failure.get("root_cause"))}
        {('<div class="failure-field full"><div class="label">Backend Evidence</div>' + backend_html + '</div>') if backend_html else ''}
        {supporting_html}
        {_failure_text_block("Business Impact", failure.get("business_impact"))}
        {_failure_text_block("Recommendation", failure.get("recommendation"))}
        {_html_readable_technical_details(dev_notes, failure)}
    </article>'''


def _failure_text_block(title: str, value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    return f'<div class="failure-field full"><div class="label">{e(title)}</div><div class="value">{e(_component_value(value))}</div></div>'


def _html_failure_conversation_rows(rows: Any) -> str:
    rows = _list(rows)
    if len(rows) <= 1:
        return ""
    blocks = []
    for row in rows[1:]:
        if not isinstance(row, list) or len(row) < 3:
            continue
        is_failure = "failure" in str(row[0]).lower()
        marker = '<div class="failure-marker">Failure occurred here</div>' if is_failure else ""
        blocks.append(
            f'<div class="failure-snippet {"failure" if is_failure else ""}>'
            f'<div class="speaker">{e(row[0])} · {e(row[1])}</div>'
            f'{marker}'
            f'<blockquote>{e(row[2])}</blockquote></div>'
        )
    return '<div class="failure-field full"><div class="label">Conversation Evidence</div>' + "".join(blocks) + "</div>"


def _html_failure_card(failure: dict[str, Any], index: int) -> str:
    metric = failure.get("metric_name") or f"Failure {index}"
    meta = (
        f'<span class="failure-badge status">{e(failure.get("status") or "FAILED")}</span>'
        f'<span class="failure-badge severity">{e(failure.get("severity") or "N/A")}</span>'
        f'<span class="failure-badge category">{e(failure.get("category") or "Uncategorized")}</span>'
    )
    field_rows = [
        ("Reason", failure.get("reason")),
        ("Expected", failure.get("expected")),
        ("Actual", failure.get("actual")),
        ("Root Cause", failure.get("root_cause")),
        ("Recommendation", failure.get("recommendation")),
    ]
    fields = "".join(
        f'<div class="failure-field {css}"><div class="label">{e(label)}</div><div class="value">{e(_component_value(value))}</div></div>'
        for label, value in field_rows
        if value not in (None, "", [], {})
        for css in ("full",)
    )
    details = _html_failure_evidence(failure.get("evidence")) + _html_failure_technical_details(failure.get("technical_details"), failure)
    snippet = _html_failure_conversation_snippet(failure)
    source = failure.get("source")
    source_html = f'<p class="muted">Source: {e(source)}</p>' if source else ""
    return f'''<article class="failure-card">
        <div class="failure-header">
            <div>
                <div class="failure-title">{e(metric)}</div>
                {source_html}
            </div>
            <div class="failure-badges">{meta}</div>
        </div>
        <div class="failure-grid">{fields or '<div class="failure-field full"><div class="value">N/A</div></div>'}</div>
        {snippet}
        {details}
    </article>'''


def _html_failure_conversation_snippet(failure: dict[str, Any]) -> str:
    snippet = [item for item in _list(failure.get("conversation_snippet")) if isinstance(item, dict)]
    if not snippet:
        return ""
    failure_turn = failure.get("conversation_turn")
    blocks = []
    for item in snippet:
        label = _value(item.get("label") or item.get("role") or "")
        turn = item.get("turn")
        is_failure = "failure" in label.lower() or (failure_turn is not None and str(turn) == str(failure_turn))
        marker = '<div class="failure-marker">Failure occurred here</div>' if is_failure else ""
        speaker = item.get("speaker") or item.get("role") or "Turn"
        message = item.get("message") or item.get("content") or item.get("text") or item
        blocks.append(
            f'<div class="failure-snippet {"failure" if is_failure else ""}">'
            f'<div class="speaker">{e(label or f"Turn {turn}")} · {e(speaker)}</div>'
            f'{marker}<blockquote>{e(_component_value(message))}</blockquote></div>'
        )
    return '<div class="failure-field full"><div class="label">Conversation Evidence</div>' + "".join(blocks) + "</div>"


def _html_failure_evidence(value: Any) -> str:
    evidence = _list(value)
    if not evidence:
        return ""
    parts = []
    for index, item in enumerate(evidence, start=1):
        parts.append(_html_readable_evidence(item, index))
    return "".join(parts)


def _html_failure_technical_details(value: Any, failure: dict[str, Any] | None = None) -> str:
    if value in (None, "", [], {}):
        return ""
    return _html_readable_technical_details(value, failure)


def _html_readable_evidence(item: Any, index: int = 1) -> str:
    if isinstance(item, dict) and str(item.get("type") or "").lower() == "conversation_card":
        return _html_conversation_card_evidence(item)
    if isinstance(item, dict):
        facts = _readable_evidence_facts(item)
        if not facts:
            return ""
        rows = "".join(f"<tr><td>{e(label)}</td><td>{e(value)}</td></tr>" for label, value in facts)
        return f'<details class="failure-details" open><summary>Supporting Evidence {index}</summary><table><tbody>{rows}</tbody></table></details>'
    return f'<details class="failure-details" open><summary>Supporting Evidence {index}</summary><p>{e(_component_value(item))}</p></details>'


def _html_conversation_card_evidence(item: dict[str, Any]) -> str:
    def block(title: str, value: Any, failed: bool = False) -> str:
        if value in (None, "", [], {}):
            value = "N/A"
        marker = '<div class="failure-marker">Failure occurred here</div>' if failed else ""
        return (
            f'<div class="failure-snippet {"failure" if failed else ""}">'
            f'<div class="speaker">{e(title)}</div>'
            f'{marker}'
            f'<blockquote>{e(_component_value(value))}</blockquote></div>'
        )

    return (
        '<div class="failure-field full"><div class="label">Conversation Evidence</div>'
        + block("Previous Turn", item.get("previous_turn"))
        + block("Failed Turn", item.get("failed_turn"), True)
        + block("Next Turn", item.get("next_turn"))
        + _failure_text_block("Why this proves the failure", item.get("why_this_proves_failure"))
        + "</div>"
    )


def _html_readable_technical_details(value: Any, failure: dict[str, Any] | None = None) -> str:
    details = value if isinstance(value, dict) else {}
    failure = failure or {}
    rows = _readable_technical_rows(details, failure)
    supporting = _supporting_evidence_from_details(details)
    if not rows and not supporting:
        return ""
    body = "".join(f"<tr><td>{e(label)}</td><td>{e(val)}</td></tr>" for label, val in rows)
    html = '<details class="failure-details" open><summary>Technical Details</summary>'
    if body:
        html += f"<table><tbody>{body}</tbody></table>"
    if supporting:
        html += '<div class="failure-field full"><div class="label">Supporting Evidence</div>' + supporting + "</div>"
    return html + "</details>"


def _readable_technical_rows(details: dict[str, Any], failure: dict[str, Any] | None = None) -> list[tuple[str, str]]:
    failure = failure or {}
    category = details.get("category") or failure.get("category")
    source = failure.get("source") or _evidence_source_label(details.get("evidence_source"), category)
    score = details.get("score", details.get("metric_score"))
    threshold = details.get("threshold")
    rows = [
        ("Metric Category", _value(category)),
        ("Evaluator", _value(source)),
        ("Score", _score_threshold_text(score, threshold)),
        ("Threshold", _percent(threshold) if threshold is not None else "N/A"),
        ("Decision", _value(failure.get("status") or "FAIL")),
        ("Conversation Position", f"Turn {failure.get('conversation_turn')}" if failure.get("conversation_turn") is not None else "N/A"),
        ("Evidence Source", _evidence_source_label(details.get("evidence_source"), category)),
    ]
    return [(label, value) for label, value in rows if value not in ("", "N/A", "None")]


def _score_threshold_text(score: Any, threshold: Any) -> str:
    if score is None:
        return "N/A"
    score_text = f"{_float(score) or 0:.2f} / 1.00"
    if threshold is not None:
        score_text += f" (threshold {_float(threshold) or 0:.2f})"
    return score_text


def _evidence_source_label(source: Any, category: Any = None) -> str:
    text = str(source or category or "").lower()
    if "backend" in text:
        return "Backend Validation"
    if "workflow" in text:
        return "Workflow Validation"
    if "rag" in text:
        return "RAG Validation"
    if "ui" in text:
        return "UI Validation"
    return "Conversation Analysis"


def _supporting_evidence_from_details(details: dict[str, Any]) -> str:
    nested = details.get("technical_details") if isinstance(details.get("technical_details"), dict) else {}
    raw = []
    raw.extend(_list(details.get("raw_metric_evidence")))
    raw.extend(_list(details.get("raw_evidence")))
    raw.extend(_list(nested.get("raw_metric_evidence")))
    components: list[dict[str, Any]] = []
    for item in _list(raw):
        if isinstance(item, dict):
            components.extend([comp for comp in _list(item.get("components")) if isinstance(comp, dict)])
            metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
            components.extend([comp for comp in _list(metadata.get("components")) if isinstance(comp, dict)])
    if not components:
        return ""
    options = [item for item in components if str(item.get("type") or item.get("subtype") or "").lower() in {"quick_reply", "suggested_prompt"}]
    other = [item for item in components if item not in options]
    parts = []
    if other:
        parts.append("<h4>UI Components Observed</h4><ul>" + "".join(f"<li>{e(_component_title(item))}</li>" for item in other) + "</ul>")
    if options:
        parts.append("<h4>Quick Reply Buttons</h4><ul>" + "".join(f"<li>{e(_option_value(item))}</li>" for item in options) + "</ul>")
    return "".join(parts)


def _supporting_evidence_rows_from_details(details: dict[str, Any]) -> list[list[str]]:
    nested = details.get("technical_details") if isinstance(details.get("technical_details"), dict) else {}
    raw = []
    raw.extend(_list(details.get("raw_metric_evidence")))
    raw.extend(_list(details.get("raw_evidence")))
    raw.extend(_list(nested.get("raw_metric_evidence")))
    components: list[dict[str, Any]] = []
    for item in _list(raw):
        if isinstance(item, dict):
            components.extend([comp for comp in _list(item.get("components")) if isinstance(comp, dict)])
            metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
            components.extend([comp for comp in _list(metadata.get("components")) if isinstance(comp, dict)])
    if not components:
        return []
    rows = [["Type", "Value"]]
    options = [item for item in components if str(item.get("type") or item.get("subtype") or "").lower() in {"quick_reply", "suggested_prompt"}]
    other = [item for item in components if item not in options]
    for item in other:
        rows.append(["UI Component", _component_title(item)])
    for item in options:
        rows.append(["Quick Reply Button", _option_value(item)])
    return rows


def _component_title(item: dict[str, Any]) -> str:
    return _value(item.get("title") or item.get("label") or item.get("name") or item.get("text") or item.get("type") or "Component")


def _readable_evidence_facts(item: dict[str, Any]) -> list[tuple[str, str]]:
    skip = {"message_id", "conversation_message_id", "conversation_message_refs", "raw_metric_evidence", "raw_conversation_snippet", "raw_json", "id"}
    facts: list[tuple[str, str]] = []
    for key, value in item.items():
        if key in skip or key.endswith("_id") or value in (None, "", [], {}):
            continue
        label = key.replace("_", " ").title()
        facts.append((label, _component_value(value)))
    return facts


def _backend_validation_rows(validation: Any) -> list[list[Any]]:
    validation = validation if isinstance(validation, dict) else {}
    rows: list[list[Any]] = [
        ["Field", "Value"],
        ["Status", validation.get("status") or "N/A"],
        ["Passed", _yes_no(validation.get("passed")) if validation.get("passed") is not None else "N/A"],
        ["Reason", validation.get("reason") or "N/A"],
    ]
    checks = validation.get("checks") if isinstance(validation.get("checks"), list) else []
    events = _backend_validation_timeline_events(validation)
    if events:
        rows.append(["Timeline", ""])
        for event in events:
            rows.append([event["title"], f"{event['status']} - {event['detail']}"])
    if checks:
        rows.append(["Check", "Result"])
        for check in checks:
            if not isinstance(check, dict):
                continue
            rows.append([
                check.get("name") or "Check",
                f"{'PASS' if check.get('passed') else 'FAIL'} - {check.get('reason') or 'N/A'}",
            ])
            evidence = check.get("evidence") if isinstance(check.get("evidence"), dict) else {}
            if evidence.get("method") or evidence.get("endpoint"):
                rows.append(["Endpoint", f"{evidence.get('method') or 'GET'} {evidence.get('endpoint') or 'N/A'}"])
            backend_connection = _backend_connection_label(evidence.get("backend_connection"))
            if backend_connection != "N/A":
                rows.append(["Backend Connection Used", backend_connection])
            if evidence.get("response_time_ms") is not None:
                rows.append(["Latency", f"{evidence.get('response_time_ms')} ms"])
    return rows


def _backend_connection_label(value: Any) -> str:
    if not isinstance(value, dict) or not value:
        return "N/A"
    name = value.get("name") or value.get("id") or "Backend connection"
    conn_type = value.get("type")
    return f"{name} ({conn_type})" if conn_type else str(name)


def _html_timeline(report: dict[str, Any]) -> str:
    turns = "".join(_html_timeline_turn(turn) for turn in report["conversation_timeline"])
    return f'<section class="section"><h2>Section 4 · Conversation Timeline</h2><div class="timeline">{turns or "N/A"}</div></section>'


def _html_timeline_turn(turn: dict[str, Any]) -> str:
    evidence = turn.get("interaction_evidence") if isinstance(turn.get("interaction_evidence"), dict) else None
    evidence_html = _html_interaction_panel(evidence) if evidence else ""
    failed_class = " failed-turn" if turn.get("failed") else ""
    fail_badge = '<span class="failure-badge status">FAIL</span>' if turn.get("failed") else ""
    return (
        f'<div class="turn{failed_class}"><div class="speaker">Turn {turn["turn"]} · {e(turn["speaker"])} {fail_badge}</div>'
        f'<div class="component-title">Response</div><blockquote>{e(turn["message"])}</blockquote>'
        f'{_html_turn_components(turn.get("components"))}'
        f'{_html_turn_attachments(turn.get("attachments"))}'
        f'{evidence_html}</div>'
    )


def _html_interaction_card(turn: dict[str, Any], evidence: dict[str, Any]) -> str:
    return f'<div class="turn interaction-turn"><div class="speaker">Turn {turn["turn"]} · EvalForge Action</div>{_html_interaction_panel(evidence)}</div>'


def _html_interaction_panel(evidence: dict[str, Any]) -> str:
    title, subtitle = _interaction_title(evidence)
    status = _value(evidence.get("status") or ("FAILED" if evidence.get("success") is False else "SUCCESS"))
    rows = _interaction_detail_rows(evidence)
    details = "".join(f"<tr><td>{e(label)}</td><td>{e(value)}</td></tr>" for label, value in rows)
    replay = _html_interaction_replay(evidence)
    media = _html_evidence_media(evidence)
    return (
        f'<div class="component-panel interaction-panel"><div class="component-title">{e(title)} <span class="status pass">{e(status)}</span></div>'
        f'<p class="muted">EvalForge Action</p>'
        f'<p class="muted">{e(subtitle)}</p>'
        f'{replay}{media}<table><tbody>{details}</tbody></table></div>'
    )


def _html_interaction_replay(evidence: dict[str, Any]) -> str:
    type_name = str(evidence.get("interaction_type") or "").lower()
    selected = evidence.get("selected_value")
    if type_name in {"form_fill", "form", "form_field_entry"} and isinstance(selected, dict):
        rows = "".join(
            f"<tr><td>{e(str(key))}</td><td>{e(_mask_sensitive(str(key), value))}</td></tr>"
            for key, value in selected.items()
        )
        return f'<table><tbody>{rows}<tr><td>Status</td><td>Submitted</td></tr></tbody></table>'
    if type_name in {"otp", "otp_entry"}:
        return '<table><tbody><tr><td>OTP</td><td>******</td></tr><tr><td>Status</td><td>Verified</td></tr></tbody></table>'
    label = "Selected"
    if type_name in {"image_upload", "camera_capture"}:
        label = "Uploaded Image" if type_name == "image_upload" else "Captured Live"
    elif type_name in {"file_upload", "document_upload", "attachment_upload"}:
        label = "Uploaded"
    elif type_name == "video_upload":
        label = "Uploaded Video"
    elif type_name == "date_picker":
        label = "Date"
    elif type_name == "time_picker":
        label = "Time"
    value = selected or evidence.get("filename") or evidence.get("artifact_path") or evidence.get("action")
    if value in (None, "", [], {}):
        return ""
    return f'<table><tbody><tr><td>{e(label)}</td><td>{e(_component_value(value))}</td></tr></tbody></table>'


def _html_artifacts(report: dict[str, Any]) -> str:
    artifacts = report["artifacts"]
    cards = "".join(
        _meta(name.replace("_", " ").title(), value if isinstance(value, int) else len(_list(value)))
        for name, value in artifacts.items()
    )
    return f'<section class="section"><h2>Artifacts</h2><div class="artifact-grid">{cards}</div></section>'


def _html_turn_components(components: Any) -> str:
    items = [item for item in _list(components) if isinstance(item, dict)]
    if not items:
        return ""
    grouped = _group_components(items)
    parts: list[str] = []
    if grouped["options"]:
        rows = "".join(
            f"<tr><td>{e(_component_type_label(item))}</td><td>{e(_option_value(item))}</td></tr>"
            for item in grouped["options"]
        )
        parts.append(f'<div class="component-panel"><div class="component-title">Available Options</div><table><thead><tr><th>Type</th><th>Value</th></tr></thead><tbody>{rows}</tbody></table></div>')
    parts.extend(_html_component(item) for item in grouped["other"])
    return '<div class="components">' + "".join(parts) + "</div>"


def _html_turn_attachments(attachments: Any) -> str:
    items = _list(attachments)
    if not items:
        return ""
    return '<div class="components"><div class="component-panel"><div class="component-title">Attachments</div>' + "".join(_html_attachment(item) for item in items) + "</div></div>"


def _html_component(component: dict[str, Any]) -> str:
    type_name = str(component.get("type") or component.get("subtype") or "").lower()
    if type_name in {"button", "quick_reply", "option", "suggested_prompt", "suggested-prompt", "link"}:
        return f'<div class="component-panel"><div class="component-title">{e(_component_type_label(component))}</div><span class="button-pill">{e(_option_value(component))}</span></div>'
    if type_name == "table":
        return _html_table_component(component)
    if type_name in {"form", "card"}:
        title = component.get("title") or component.get("label") or type_name.title()
        fields = component.get("fields") or component.get("items") or component.get("sections") or []
        body = _html_fields(fields) if _list(fields) else _html_key_value_rows(component)
        return f'<div class="component-panel"><div class="component-title">{e(title)}</div>{body}</div>'
    if type_name == "image":
        url = _safe_url(component.get("url") or component.get("src") or component.get("path"))
        caption = e(component.get("name") or component.get("alt") or "Image")
        if url:
            return f'<div class="component-panel"><div class="component-title">{caption}</div><img class="rich-image" src="{e(url)}" alt="{caption}"></div>'
    if type_name == "video":
        url = _safe_url(component.get("url") or component.get("src") or component.get("path"))
        caption = e(component.get("name") or component.get("alt") or "Video")
        if url:
            return f'<div class="component-panel"><div class="component-title">{caption}</div><video class="rich-video" controls src="{e(url)}"></video></div>'
    if type_name in {"attachment", "file"}:
        return _html_attachment(component)
    if _looks_like_otp(component):
        return f'<div class="component-panel"><div class="component-title">Verification</div><table><tbody><tr><td>OTP Requested</td><td>{e(_yes_no(component.get("requested", True)))}</td></tr><tr><td>OTP Verified</td><td>{e(_yes_no(component.get("verified") or component.get("success")))}</td></tr></tbody></table></div>'
    return _html_key_value_component(type_name or "Component", component)


def _html_table_component(component: dict[str, Any]) -> str:
    table_data = component.get("data") if isinstance(component.get("data"), dict) else component
    if not _list(table_data.get("rows")) and isinstance(table_data.get("sections"), list):
        rendered = "".join(_html_table_component(section) for section in table_data["sections"] if isinstance(section, dict))
        return rendered or _html_key_value_component("Table", component)
    table_rows = _component_table_rows(component)
    if not table_rows:
        return _html_key_value_component("Table", component)
    headers = table_rows[0]
    rows = table_rows[1:]
    head = "".join(f"<th>{e(header)}</th>" for header in headers)
    body = "".join(
        "<tr>" + "".join(f"<td>{e(cell)}</td>" for cell in row) + "</tr>"
        for row in rows[:100]
    )
    return f'<div class="component-panel"><div class="component-title">{e(component.get("title") or "Table")}</div><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'


def _html_fields(fields: Any) -> str:
    rows = [
        f"<tr><td>{e(row[0])}</td><td>{e(row[1])}</td></tr>"
        for row in _component_field_rows(fields)
    ]
    return f'<table><tbody>{"".join(rows)}</tbody></table>' if rows else ""


def _html_key_value_component(title: str, component: dict[str, Any]) -> str:
    return f'<div class="component-panel"><div class="component-title">{e(title.replace("_", " ").title())}</div>{_html_key_value_rows(component)}</div>'


def _html_key_value_rows(component: dict[str, Any]) -> str:
    ignored = {"type", "subtype", "metadata", "provider_metadata"}
    rows = []
    for key, value in component.items():
        if key in ignored or value in (None, "", [], {}):
            continue
        rows.append(f"<tr><td>{e(str(key).replace('_', ' ').title())}</td><td>{e(_component_value(value))}</td></tr>")
    return f'<table><tbody>{"".join(rows)}</tbody></table>' if rows else '<p class="muted">N/A</p>'


def _html_attachment(item: Any) -> str:
    component = item if isinstance(item, dict) else {"path": item, "name": item}
    url = _safe_url(component.get("preview_url") or component.get("thumbnail_url") or component.get("url") or component.get("path") or component.get("src") or component.get("blob_url"))
    label = e(_attachment_label(component))
    meta = e(component.get("mime_type") or component.get("type") or "Attachment")
    if _is_image_attachment(component) and url:
        return f'<div class="component-panel"><div class="component-title">Uploaded Image</div><img class="rich-image" src="{e(url)}" alt="{label}"><div class="muted">{label} · {meta}</div></div>'
    link = f'<a href="{e(url)}" target="_blank" rel="noopener">{label}</a>' if url else label
    return f'<div class="attachment-card"><strong>{link}</strong><div class="muted">{meta}</div></div>'


def _html_evidence_media(evidence: dict[str, Any]) -> str:
    path = _evidence_image_path(evidence)
    if not path:
        return ""
    label = e(evidence.get("filename") or os.path.basename(path))
    return f'<div class="component-panel"><div class="component-title">Uploaded Image</div><img class="rich-image" src="{e(path)}" alt="{label}"><div class="muted">{label}</div></div>'


def _html_risks(report: dict[str, Any]) -> str:
    items = "".join(f'<div class="risk"><strong>{e(item["level"])} Risk</strong><p>{e(item["risk"])}</p><div class="muted">{e(item["impact"])}</div></div>' for item in report["risks"]) or '<p class="muted">N/A</p>'
    return f'<section class="section"><h2>Risks</h2>{items}</section>'


def _html_recommendations(report: dict[str, Any]) -> str:
    items = "".join(f'<div class="recommendation">{e(item)}</div>' for item in report["recommendations"])
    return f'<section class="section"><h2>Recommendations</h2>{items}</section>'


def _html_json_section(title: str, value: Any) -> str:
    if value is None or value == [] or value == {} or value == "":
        value = "N/A"
    if title.lower() == "appendix":
        return _html_appendix_summary(value)
    return f'<section class="section"><h2>{e(title)}</h2>{_json_dump(value)}</section>'


def _html_appendix_summary(value: Any) -> str:
    appendix = value if isinstance(value, dict) else {}
    rows = _appendix_summary_rows(appendix)
    body = "".join(f"<tr><td>{e(row[0])}</td><td>{e(row[1])}</td></tr>" for row in rows[1:])
    return f'<section class="section"><h2>Appendix</h2><p class="muted">Raw debugging payloads are available in the saved evaluation JSON and are intentionally not rendered in the report.</p><table><tbody>{body}</tbody></table></section>'


def _json_dump(value: Any) -> str:
    return f"<pre>{e(json.dumps(value, indent=2, default=str))}</pre>"


def _meta(label: str, value: Any) -> str:
    return f'<div class="meta"><div class="label">{e(label)}</div><div class="value">{e(value)}</div></div>'


def _reportlab_pdf(report: dict[str, Any], theme: ReportTheme) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise RuntimeError("PDF export requires reportlab. Install dependencies from requirements.txt.") from exc

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontSize=28, leading=34, alignment=TA_CENTER, textColor=colors.HexColor(theme.ink), spaceAfter=18))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontSize=15, leading=19, textColor=colors.HexColor(theme.primary), spaceBefore=14, spaceAfter=9))
    styles.add(ParagraphStyle(name="SmallMuted", parent=styles["BodyText"], fontSize=8.5, leading=11, textColor=colors.HexColor(theme.muted)))
    styles.add(ParagraphStyle(name="BubbleUser", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#eef2ff"), borderColor=colors.HexColor("#c7d2fe"), borderWidth=0.5, borderPadding=7, spaceAfter=6))
    styles.add(ParagraphStyle(name="BubbleAgent", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#f8fafc"), borderColor=colors.HexColor("#e2e8f0"), borderWidth=0.5, borderPadding=7, spaceAfter=6))
    styles.add(ParagraphStyle(name="FailureBubble", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#fef2f2"), borderColor=colors.HexColor("#ef4444"), borderWidth=0.75, borderPadding=7, spaceAfter=6))
    styles.add(ParagraphStyle(name="CodeBlock", parent=styles["Code"], fontName="Courier", fontSize=7, leading=8.5, backColor=colors.HexColor("#f8fafc"), borderColor=colors.HexColor("#dbe3ef"), borderWidth=0.35, borderPadding=6, spaceAfter=6))
    story: list[Any] = []
    h = report["header"]

    story.extend([
        Spacer(1, 0.9 * inch),
        Paragraph(theme.brand.upper(), styles["SmallMuted"]),
        Paragraph(report.get("title") or theme.title, styles["CoverTitle"]),
        Paragraph(_xml(report["executive_summary"]["narrative"]), styles["BodyText"]),
        Spacer(1, 0.25 * inch),
        *(_pdf_scorecards(report, theme)),
        Spacer(1, 0.2 * inch),
        _pdf_kv_table([
            ("Evaluation ID", h["evaluation_id"]),
            ("Agent", h["agent_name"]),
            ("Scenario", h["scenario"]),
            ("Adapter", h["adapter"]),
            ("Date & Time", h["date_time"]),
            ("Status", h["status"]),
            ("Client Report Version", report["report_metadata"].get("client_report_version")),
            ("Generated By", report["report_metadata"].get("generated_by")),
            ("EvalForge Version", report["report_metadata"].get("evalforge_version")),
            ("Generated At", report["report_metadata"].get("generated_at")),
        ], theme),
        PageBreak(),
    ])

    story.append(Paragraph("Executive Summary", styles["SectionTitle"]))
    story.append(_pdf_kv_table([
        ("Project", report["executive_summary"].get("project")),
        ("Overall Risk", report["executive_summary"].get("overall_risk")),
        ("Total Metrics", report["executive_summary"].get("total_metrics")),
        ("Passed", report["executive_summary"].get("passed")),
        ("Failed", report["executive_summary"].get("failed")),
        ("Critical Failures", report["executive_summary"].get("critical_failures")),
        ("Recommendation", report["executive_summary"].get("recommendation")),
    ], theme))
    story.append(Paragraph("Evaluation Summary", styles["SectionTitle"]))
    story.append(_pdf_table(
        [["Metric", "Category", "Score", "Threshold", "Status", "Severity"]]
        + [[item["metric"], item["category"], item["score_display"], item["threshold_display"], item["status"], item["severity"]] for item in report.get("evaluation_summary") or []],
        theme,
    ))
    story.append(Paragraph("Executive Dashboard", styles["SectionTitle"]))
    story.append(_pdf_metric_chart(report, theme))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Capability Matrix", styles["SectionTitle"]))
    story.append(_pdf_table([["Capability", "Status", "Evidence"]] + [[r["capability"], r["status"], r["evidence"]] for r in report["capability_matrix"]], theme))

    story.append(Paragraph("Conversation Timeline", styles["SectionTitle"]))
    for turn in report["conversation_timeline"][:40]:
        evidence = turn.get("interaction_evidence") if isinstance(turn.get("interaction_evidence"), dict) else None
        if evidence:
            title, subtitle = _interaction_title(evidence)
            story.append(Paragraph(f"<b>Turn {turn['turn']} · EvalForge Action · {_xml(title)}</b><br/>{_xml(subtitle)}", styles["BubbleAgent"]))
            _pdf_evidence_image(story, evidence)
            rows = [["Field", "Value"]] + _interaction_detail_rows(evidence)
            story.append(_pdf_table(rows, theme))
            continue
        style = styles["BubbleUser"] if str(turn["speaker"]).lower() == "user" else styles["BubbleAgent"]
        story.append(Paragraph(f"<b>Turn {turn['turn']} · {_xml(turn['speaker'])}</b><br/>{_xml(turn['message'])}", style))
        _pdf_turn_components(story, turn, styles, theme)

    story.append(Paragraph("AI Findings", styles["SectionTitle"]))
    story.append(_pdf_table([["Finding", "Severity", "Score", "Summary"]] + [[f["title"], f["severity"], f["score"], f["summary"]] for f in report["ai_findings"]], theme))
    story.append(Paragraph("Risk Analysis", styles["SectionTitle"]))
    story.append(_pdf_table([["Level", "Risk", "Impact"]] + [[r["level"], r["risk"], r["impact"]] for r in (report["risks"] or [{"level": "N/A", "risk": "N/A", "impact": "N/A"}])], theme))
    story.append(Paragraph("Backend Validation", styles["SectionTitle"]))
    story.append(_pdf_table(_backend_validation_rows(report.get("backend_validation")), theme))
    if report.get("workflow_validation"):
        story.append(Paragraph("Workflow Validation", styles["SectionTitle"]))
        story.append(_pdf_table(
            [["Step", "Expected", "Observed", "Status"]]
            + [[row.get("step"), row.get("expected"), row.get("observed"), row.get("status")] for row in report.get("workflow_validation")],
            theme,
        ))
    if report.get("rag_validation") and (report["rag_validation"].get("metrics") or report["rag_validation"].get("questions")):
        story.append(Paragraph("RAG Validation", styles["SectionTitle"]))
        story.append(_pdf_table(
            [["Metric", "Score", "Status", "Evidence"]]
            + [[m.get("name"), _percent(m.get("score")), m.get("status"), m.get("reason")] for m in report["rag_validation"].get("metrics") or []],
            theme,
        ))
    _pdf_failure_analysis(story, report, styles, theme)
    story.append(Paragraph("Overall Findings", styles["SectionTitle"]))
    story.append(_pdf_table(
        [["Category", "Issues", "Metrics Evaluated"]]
        + [[item.get("category"), item.get("issues"), item.get("metrics")] for item in report.get("overall_findings") or []],
        theme,
    ))
    story.append(Paragraph("Priority Recommendations", styles["SectionTitle"]))
    for item in report.get("priority_recommendations") or []:
        story.append(Paragraph(f"Priority {item.get('priority')}: {_xml(item.get('title'))} - {_xml(item.get('recommendation'))}", styles["BodyText"]))
    story.append(Paragraph("Artifact Summary", styles["SectionTitle"]))
    story.append(_pdf_table([["Artifact", "Count"]] + [[k.replace("_", " ").title(), str(_artifact_count(v))] for k, v in report["artifacts"].items()], theme))
    story.append(Paragraph("Appendix", styles["SectionTitle"]))
    story.append(_pdf_table(_appendix_summary_rows(report.get("appendix")), theme))

    doc.build(story)
    return output.getvalue()


def _reportlab_executive_pdf(report: dict[str, Any], theme: ReportTheme) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("PDF export requires reportlab. Install dependencies from requirements.txt.") from exc

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontSize=28, leading=34, alignment=TA_CENTER, textColor=colors.HexColor(theme.ink), spaceAfter=18))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontSize=15, leading=19, textColor=colors.HexColor(theme.primary), spaceBefore=14, spaceAfter=9))
    styles.add(ParagraphStyle(name="SmallMuted", parent=styles["BodyText"], fontSize=8.5, leading=11, textColor=colors.HexColor(theme.muted)))
    styles.add(ParagraphStyle(name="BubbleUser", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#eef2ff"), borderColor=colors.HexColor("#c7d2fe"), borderWidth=0.5, borderPadding=7, spaceAfter=6))
    styles.add(ParagraphStyle(name="BubbleAgent", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#f8fafc"), borderColor=colors.HexColor("#e2e8f0"), borderWidth=0.5, borderPadding=7, spaceAfter=6))
    styles.add(ParagraphStyle(name="FailureBubble", parent=styles["BodyText"], fontSize=9, leading=12, backColor=colors.HexColor("#fef2f2"), borderColor=colors.HexColor("#ef4444"), borderWidth=0.75, borderPadding=7, spaceAfter=6))

    h = report["header"]
    story: list[Any] = [
        Spacer(1, 0.85 * inch),
        Paragraph(theme.brand.upper(), styles["SmallMuted"]),
        Paragraph(report["title"], styles["CoverTitle"]),
        Paragraph(_xml(report["executive_summary"]["narrative"]), styles["BodyText"]),
        Spacer(1, 0.22 * inch),
        _pdf_kv_table([
            ("Client", report["executive_summary"].get("project")),
            ("Agent", h.get("agent_name")),
            ("Scenario", h.get("scenario")),
            ("Evaluation Date", h.get("date_time")),
            ("Overall Score", _percent(h.get("overall_score"))),
            ("Overall Risk", report["executive_summary"].get("overall_risk")),
            ("Evaluation ID", h.get("evaluation_id")),
            ("Version", report["report_metadata"].get("evalforge_version")),
            ("Generated By", report["report_metadata"].get("generated_by")),
        ], theme),
        PageBreak(),
    ]
    story.append(Paragraph("Executive Summary", styles["SectionTitle"]))
    story.append(Paragraph(_xml(report["executive_summary"].get("narrative")), styles["BodyText"]))
    story.append(_pdf_kv_table([
        ("Overall Performance", report["executive_summary"].get("overall_score_display")),
        ("Major Strengths", _executive_strengths(report)),
        ("Major Weaknesses", _executive_weaknesses(report)),
        ("Production Readiness", report["executive_summary"].get("recommendation")),
        ("Recommendation", _final_recommendation_label(report)),
    ], theme))
    story.append(PageBreak())
    story.append(Paragraph("KPI Dashboard", styles["SectionTitle"]))
    story.append(_pdf_table([["KPI", "Value", "Comment"]] + [[item["label"], item["value"], item.get("note")] for item in _executive_kpis(report)], theme))
    story.append(Paragraph("Capability Matrix", styles["SectionTitle"]))
    story.append(_pdf_table([["Capability", "Status", "Evidence"]] + [[r["capability"], r["status"], r["evidence"]] for r in report["capability_matrix"]], theme))

    failures = [item for item in _list(report.get("detailed_failure_analysis")) if isinstance(item, dict)]
    for failure in failures:
        story.append(PageBreak())
        story.append(Paragraph(_xml(failure.get("metric_name") or "Failure Analysis"), styles["SectionTitle"]))
        story.append(_pdf_kv_table([
            ("Severity", failure.get("severity")),
            ("Category", failure.get("category")),
            ("Business Summary", failure.get("business_summary")),
            ("Expected Behaviour", failure.get("expected_behaviour")),
            ("Actual Behaviour", failure.get("actual_behaviour")),
            ("Business Impact", failure.get("business_impact")),
            ("Recommendation", failure.get("recommendation")),
        ], theme))
        _pdf_failure_conversation_from_rows(story, failure.get("conversation_evidence"), styles)

    if report.get("backend_validation"):
        story.append(PageBreak())
        story.append(Paragraph("Backend Validation", styles["SectionTitle"]))
        story.append(Paragraph("Backend validation compares the agent outcome against configured business systems.", styles["BodyText"]))
        story.append(_pdf_table(_backend_validation_rows(report.get("backend_validation")), theme))
    if report.get("workflow_validation"):
        story.append(Paragraph("Workflow Validation", styles["SectionTitle"]))
        story.append(_pdf_table(
            [["Step", "Expected", "Observed", "Status"]]
            + [[row.get("step"), row.get("expected"), row.get("observed"), row.get("status")] for row in report.get("workflow_validation")],
            theme,
        ))
    if report.get("rag_validation") and (report["rag_validation"].get("metrics") or report["rag_validation"].get("questions")):
        story.append(Paragraph("RAG Validation", styles["SectionTitle"]))
        story.append(_pdf_table(
            [["Question / Metric", "Answer / Evidence", "Score"]]
            + [[m.get("name"), m.get("reason"), _percent(m.get("score"))] for m in report["rag_validation"].get("metrics") or []],
            theme,
        ))
    story.append(Paragraph("Overall Findings", styles["SectionTitle"]))
    story.append(Paragraph(_xml(_executive_findings_narrative(report)), styles["BodyText"]))
    story.append(Paragraph("Priority Recommendations", styles["SectionTitle"]))
    for item in report.get("priority_recommendations") or []:
        story.append(Paragraph(f"<b>Priority {item.get('priority')}: {_xml(item.get('title'))}</b><br/>{_xml(item.get('recommendation'))}", styles["BodyText"]))
    story.append(Paragraph("Final Recommendation", styles["SectionTitle"]))
    story.append(Paragraph(f"<b>{_xml(_final_recommendation_label(report))}</b>: {_xml(report['executive_summary'].get('recommendation'))}", styles["BodyText"]))
    doc.build(story)
    return output.getvalue()


def _pdf_failure_analysis(story: list[Any], report: dict[str, Any], styles: Any, theme: ReportTheme) -> None:
    from reportlab.platypus import Paragraph, Spacer

    failures = [item for item in _list(report.get("failure_analysis")) if isinstance(item, dict)]
    if not failures:
        return
    story.append(Paragraph("Failure Analysis", styles["SectionTitle"]))
    for index, failure in enumerate(failures, start=1):
        metric = failure.get("metric_name") or f"Failure {index}"
        story.append(Paragraph(f"<b>{_xml(metric)}</b>", styles["Heading3"]))
        story.append(_pdf_kv_table([
            ("Status", failure.get("status") or "FAILED"),
            ("Severity", failure.get("severity") or "N/A"),
            ("Category", failure.get("category") or "Uncategorized"),
            ("Source", failure.get("source") or "N/A"),
            ("Conversation Turn", failure.get("conversation_turn") or "N/A"),
        ], theme))
        _pdf_long_text_section(story, "Reason", failure.get("reason"), styles)
        _pdf_long_text_section(story, "Expected", failure.get("expected"), styles)
        _pdf_long_text_section(story, "Actual", failure.get("actual"), styles)
        _pdf_long_text_section(story, "Root Cause", failure.get("root_cause"), styles)
        _pdf_long_text_section(story, "Recommendation", failure.get("recommendation"), styles)
        _pdf_failure_conversation_section(story, failure, styles)
        _pdf_failure_evidence_section(story, failure.get("evidence"), styles, theme)
        details = failure.get("technical_details")
        if details not in (None, "", [], {}):
            _pdf_readable_technical_details(story, details, failure, styles, theme)
        story.append(Spacer(1, 10))


def _pdf_scorecards(report: dict[str, Any], theme: ReportTheme) -> list[Any]:
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    styles = getSampleStyleSheet()
    rows = [[
        f"Overall\n{_percent(report['header']['overall_score'])}",
        f"Status\n{report['header']['status']}",
        f"Passed Metrics\n{report['scorecard']['passed_metrics']}",
        f"Failed Metrics\n{report['scorecard']['failed_metrics']}",
    ]]
    table = Table(rows, colWidths=[1.75 * inch] * 4, rowHeights=[0.75 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor(theme.ink)),
        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#dbe3ef")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dbe3ef")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
    ]))
    
    flowables = [table, Spacer(1, 15)]
    categories = {}
    for metric in report["scorecard"]["metrics"]:
        cat = metric.get("category") or "Conversation"
        categories.setdefault(cat, []).append(metric)
        
    for cat, metrics in categories.items():
        flowables.append(Paragraph(f"{str(cat).title()} Metrics", styles["Heading3"]))
        metric_rows = [["Metric", "Score", "Status", "Reason"]]
        for m in metrics:
            metric_rows.append([m["name"], m["display"], m["status"], m["reason"]])
        flowables.append(_pdf_table(metric_rows, theme))
        flowables.append(Spacer(1, 15))
    return flowables


def _pdf_metric_chart(report: dict[str, Any], theme: ReportTheme) -> Any:
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    metrics = report["scorecard"]["metrics"][:6]
    values = [[_bar_width(m["score"]) for m in metrics] or [0]]
    drawing = Drawing(460, 170)
    chart = VerticalBarChart()
    chart.x = 35
    chart.y = 35
    chart.height = 100
    chart.width = 390
    chart.data = values
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = 100
    chart.valueAxis.valueStep = 20
    chart.categoryAxis.categoryNames = [m["name"][:12] for m in metrics] or ["N/A"]
    chart.bars[0].fillColor = colors.HexColor(theme.primary)
    drawing.add(String(35, 148, "Metric Scores", fontName="Helvetica-Bold", fontSize=11, fillColor=colors.HexColor(theme.ink)))
    drawing.add(chart)
    return drawing


def _pdf_kv_table(items: list[tuple[str, Any]], theme: ReportTheme) -> Any:
    return _pdf_table([["Field", "Value"]] + [[label, _value(value)] for label, value in items], theme)


def _pdf_table(rows: list[list[Any]], theme: ReportTheme) -> Any:
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    styles = getSampleStyleSheet()
    wrapped = [[Paragraph(_xml(_pdf_cell_value(cell)), styles["BodyText"]) for cell in row] for row in rows]
    table = Table(wrapped, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(theme.primary)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#dbe3ef")),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#ffffff")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _pdf_cell_value(value: Any) -> str:
    text = _component_value(value) if isinstance(value, (dict, list, tuple)) else _value(value)
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= PDF_TABLE_CELL_MAX_CHARS:
        return text
    suffix = " ... [truncated]"
    return text[: PDF_TABLE_CELL_MAX_CHARS - len(suffix)].rstrip() + suffix


def _pdf_long_text_section(story: list[Any], title: str, value: Any, styles: Any) -> None:
    from reportlab.platypus import Paragraph, Spacer

    if value in (None, "", [], {}):
        return
    story.append(Paragraph(f"<b>{_xml(title)}</b>", styles["SmallMuted"]))
    text = _component_value(value)
    for chunk in _split_pdf_text(text, PDF_LONG_TEXT_CHUNK_CHARS):
        story.append(Paragraph(_xml(chunk), styles["BodyText"]))
    story.append(Spacer(1, 4))


def _pdf_preformatted_section(story: list[Any], title: str, value: Any, styles: Any) -> None:
    from reportlab.platypus import Paragraph, Preformatted, Spacer

    story.append(Paragraph(f"<b>{_xml(title)}</b>", styles["SmallMuted"]))
    text = _pretty_json_text(value)
    text = _wrap_pdf_code_lines(text)
    for chunk in _split_pdf_text(text, PDF_CODE_BLOCK_CHUNK_CHARS):
        story.append(Preformatted(_xml(chunk), styles["CodeBlock"]))
    story.append(Spacer(1, 4))


def _pdf_failure_conversation_section(story: list[Any], failure: dict[str, Any], styles: Any) -> None:
    from reportlab.platypus import Paragraph, Spacer

    snippet = [item for item in _list(failure.get("conversation_snippet")) if isinstance(item, dict)]
    if not snippet:
        return
    story.append(Paragraph("<b>Conversation Evidence</b>", styles["SmallMuted"]))
    failure_turn = failure.get("conversation_turn")
    for item in snippet:
        turn = item.get("turn") or item.get("conversation_turn") or "N/A"
        label = _value(item.get("label") or "")
        speaker = item.get("speaker") or item.get("role") or "N/A"
        message = item.get("message") or item.get("content") or item.get("text") or item
        is_failure = "failure" in label.lower() or (failure_turn is not None and str(turn) == str(failure_turn))
        style = styles["FailureBubble"] if is_failure else styles["BubbleAgent"]
        marker = " - FAILURE OCCURRED HERE" if is_failure else ""
        story.append(Paragraph(f"<b>{_xml(label or f'Turn {turn}')}{_xml(marker)} · {_xml(speaker)}</b><br/>{_xml(_component_value(message))}", style))
    story.append(Spacer(1, 4))


def _pdf_failure_conversation_from_rows(story: list[Any], rows: Any, styles: Any) -> None:
    from reportlab.platypus import Paragraph, Spacer

    rows = _list(rows)
    if len(rows) <= 1:
        return
    story.append(Paragraph("<b>Conversation Evidence</b>", styles["SmallMuted"]))
    for row in rows[1:]:
        if not isinstance(row, list) or len(row) < 3:
            continue
        is_failure = "failure" in str(row[0]).lower()
        style = styles["FailureBubble"] if is_failure else styles["BubbleAgent"]
        marker = " - FAILED" if is_failure else ""
        story.append(Paragraph(f"<b>{_xml(row[0])}{marker} · {_xml(row[1])}</b><br/>{_xml(row[2])}", style))
    story.append(Spacer(1, 4))


def _pdf_failure_evidence_section(story: list[Any], value: Any, styles: Any, theme: ReportTheme) -> None:
    from reportlab.platypus import Paragraph

    evidence = _list(value)
    if not evidence:
        return
    story.append(Paragraph("<b>Evidence</b>", styles["SmallMuted"]))
    for index, item in enumerate(evidence, start=1):
        if isinstance(item, dict) and str(item.get("type") or "").lower() == "conversation_card":
            _pdf_conversation_card_evidence(story, item, styles)
            continue
        facts = _readable_evidence_facts(item) if isinstance(item, dict) else [("Evidence", _component_value(item))]
        if facts:
            story.append(Paragraph(f"<b>Supporting Evidence {index}</b>", styles["SmallMuted"]))
            story.append(_pdf_table([["Field", "Value"]] + [[label, value] for label, value in facts], theme))


def _pdf_conversation_card_evidence(story: list[Any], item: dict[str, Any], styles: Any) -> None:
    from reportlab.platypus import Paragraph

    story.append(Paragraph("<b>Conversation Evidence</b>", styles["SmallMuted"]))
    for title, key, failed in (
        ("Previous Turn", "previous_turn", False),
        ("Failed Turn", "failed_turn", True),
        ("Next Turn", "next_turn", False),
    ):
        style = styles["FailureBubble"] if failed else styles["BubbleAgent"]
        story.append(Paragraph(f"<b>{_xml(title)}</b><br/>{_xml(_component_value(item.get(key) or 'N/A'))}", style))
    _pdf_long_text_section(story, "Why this proves the failure", item.get("why_this_proves_failure"), styles)


def _pdf_readable_technical_details(story: list[Any], value: Any, failure: dict[str, Any], styles: Any, theme: ReportTheme) -> None:
    from reportlab.platypus import Paragraph

    details = value if isinstance(value, dict) else {}
    rows = _readable_technical_rows(details, failure)
    supporting = _supporting_evidence_rows_from_details(details)
    if rows:
        story.append(Paragraph("<b>Technical Details</b>", styles["SmallMuted"]))
        story.append(_pdf_table([["Field", "Value"]] + [[label, val] for label, val in rows], theme))
    if supporting:
        story.append(Paragraph("<b>Supporting Evidence</b>", styles["SmallMuted"]))
        story.append(_pdf_table(supporting, theme))


def _pretty_json_text(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, indent=2, default=str)
    return _value(value)


def _wrap_pdf_code_lines(text: str, width: int = 96) -> str:
    wrapped: list[str] = []
    for line in str(text).splitlines() or [""]:
        if len(line) <= width:
            wrapped.append(line)
        else:
            wrapped.extend(textwrap.wrap(line, width=width, replace_whitespace=False, drop_whitespace=False) or [line])
    return "\n".join(wrapped)


def _split_pdf_text(text: str, max_chars: int) -> list[str]:
    text = _value(text)
    if len(text) <= max_chars:
        return [text]
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0
    for line in text.splitlines() or [text]:
        line_len = len(line) + 1
        if current and current_len + line_len > max_chars:
            chunks.append("\n".join(current))
            current = []
            current_len = 0
        if line_len > max_chars:
            chunks.extend(textwrap.wrap(line, width=max_chars, replace_whitespace=False, drop_whitespace=False) or [line])
            continue
        current.append(line)
        current_len += line_len
    if current:
        chunks.append("\n".join(current))
    return chunks or [""]


def _appendix_summary_rows(appendix: Any) -> list[list[Any]]:
    appendix = appendix if isinstance(appendix, dict) else {}
    return [
        ["Appendix", "Details"],
        ["Conversation IDs", f"{len(_list(appendix.get('conversation_ids')))} captured"],
        ["Evidence IDs", ", ".join(str(item) for item in _list(appendix.get("evidence_ids"))) or "N/A"],
        ["API Logs", f"{len(_list(appendix.get('api_logs')))} item(s)"],
        ["Attachments", f"{len(_list(appendix.get('attachments')))} item(s)"],
        ["Screenshots", f"{len(_list(appendix.get('screenshots')))} item(s)"],
        ["Tool Calls", f"{len(_list(appendix.get('tool_calls')))} item(s)"],
        ["Backend Requests", f"{len(_list(appendix.get('backend_requests')))} item(s)"],
        ["Raw Data", "Available separately in the saved evaluation JSON."],
    ]


def _pdf_turn_components(story: list[Any], turn: dict[str, Any], styles: Any, theme: ReportTheme) -> None:
    from reportlab.platypus import Paragraph

    components = [item for item in _list(turn.get("components")) if isinstance(item, dict)]
    grouped = _group_components(components)
    if grouped["options"]:
        story.append(Paragraph("<b>Available Options</b>", styles["SmallMuted"]))
        story.append(_pdf_table(
            [["Type", "Value"]] + [[_component_type_label(item), _option_value(item)] for item in grouped["options"]],
            theme,
        ))
    for component in grouped["other"]:
        type_name = str(component.get("type") or component.get("subtype") or "").lower()
        if type_name == "table":
            rows = _component_table_rows(component)
            if rows:
                story.append(Paragraph(f"<b>{_xml(component.get('title') or 'Table')}</b>", styles["SmallMuted"]))
                story.append(_pdf_table(rows, theme))
        elif type_name in {"form", "card"}:
            fields = component.get("fields") or component.get("items") or component.get("sections") or []
            rows = [["Field", "Value"]] + _component_field_rows(fields)
            story.append(Paragraph(f"<b>{_xml(component.get('title') or component.get('label') or type_name.title())}</b>", styles["SmallMuted"]))
            story.append(_pdf_table(rows, theme))
        else:
            for line in _turn_artifact_lines({"components": [component], "attachments": []}):
                story.append(Paragraph(_xml(line), styles["SmallMuted"]))
    attachments = _list(turn.get("attachments"))
    if attachments:
        story.append(Paragraph("<b>Attachments</b>", styles["SmallMuted"]))
        for attachment in attachments:
            story.append(Paragraph(_xml(_attachment_label(attachment)), styles["SmallMuted"]))


def _python_docx_executive(report: dict[str, Any], theme: ReportTheme) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("DOCX export requires python-docx. Install dependencies from requirements.txt.") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    h = report["header"]
    brand = document.add_paragraph(theme.brand.upper())
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.runs[0].font.size = Pt(9)
    brand.runs[0].font.bold = True
    title = document.add_heading(report["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(report["executive_summary"]["narrative"])
    _docx_table(document, "Cover Metadata", [
        ["Field", "Value"],
        ["Client", report["executive_summary"].get("project")],
        ["Agent", h.get("agent_name")],
        ["Scenario", h.get("scenario")],
        ["Evaluation Date", h.get("date_time")],
        ["Overall Score", _percent(h.get("overall_score"))],
        ["Overall Risk", report["executive_summary"].get("overall_risk")],
        ["Evaluation ID", h.get("evaluation_id")],
        ["Version", report["report_metadata"].get("evalforge_version")],
        ["Generated By", report["report_metadata"].get("generated_by")],
    ])
    document.add_page_break()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(report["executive_summary"].get("narrative"))
    _docx_table(document, "KPI Dashboard", [["KPI", "Value", "Comment"]] + [[item["label"], item["value"], item.get("note")] for item in _executive_kpis(report)])
    _docx_table(document, "Capability Matrix", [["Capability", "Status", "Evidence"]] + [[r["capability"], r["status"], r["evidence"]] for r in report["capability_matrix"]])
    failures = [item for item in _list(report.get("detailed_failure_analysis")) if isinstance(item, dict)]
    for failure in failures:
        document.add_page_break()
        document.add_heading(failure.get("metric_name") or "Failure Analysis", level=1)
        _docx_table(document, "Failure Summary", [
            ["Field", "Value"],
            ["Severity", failure.get("severity")],
            ["Business Summary", failure.get("business_summary")],
            ["Expected Behaviour", failure.get("expected_behaviour")],
            ["Actual Behaviour", failure.get("actual_behaviour")],
            ["Business Impact", failure.get("business_impact")],
            ["Recommendation", failure.get("recommendation")],
        ])
        _docx_failure_conversation_rows(document, failure.get("conversation_evidence"))
    if report.get("backend_validation"):
        _docx_table(document, "Backend Validation", _backend_validation_rows(report.get("backend_validation")))
    if report.get("workflow_validation"):
        _docx_table(
            document,
            "Workflow Validation",
            [["Step", "Expected", "Observed", "Status"]]
            + [[row.get("step"), row.get("expected"), row.get("observed"), row.get("status")] for row in report.get("workflow_validation")],
        )
    document.add_heading("Overall Findings", level=1)
    document.add_paragraph(_executive_findings_narrative(report))
    document.add_heading("Priority Recommendations", level=1)
    for item in report.get("priority_recommendations") or []:
        p = document.add_paragraph()
        p.add_run(f"Priority {item.get('priority')}: {item.get('title')}: ").bold = True
        p.add_run(_value(item.get("recommendation")))
    document.add_heading("Final Recommendation", level=1)
    document.add_paragraph(f"{_final_recommendation_label(report)}: {report['executive_summary'].get('recommendation')}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _python_docx(report: dict[str, Any], theme: ReportTheme) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("DOCX export requires python-docx. Install dependencies from requirements.txt.") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    h = report["header"]
    brand = document.add_paragraph(theme.brand.upper())
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.runs[0].font.size = Pt(9)
    brand.runs[0].font.bold = True
    brand.runs[0].font.color.rgb = _rgb(theme.primary)
    title = document.add_heading(report.get("title") or theme.title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(report["executive_summary"]["narrative"])
    _docx_scorecards(document, report)
    _docx_table(document, "Assessment Metadata", [
        ["Field", "Value"],
        ["Evaluation ID", h["evaluation_id"]],
        ["Agent", h["agent_name"]],
        ["Scenario", h["scenario"]],
        ["Adapter", h["adapter"]],
        ["Date & Time", h["date_time"]],
        ["Status", h["status"]],
        ["Client Report Version", report["report_metadata"].get("client_report_version")],
        ["Generated By", report["report_metadata"].get("generated_by")],
        ["EvalForge Version", report["report_metadata"].get("evalforge_version")],
        ["Generated At", report["report_metadata"].get("generated_at")],
    ])
    document.add_page_break()
    _docx_table(document, "Executive Summary", [
        ["Field", "Value"],
        ["Project", report["executive_summary"].get("project")],
        ["Overall Risk", report["executive_summary"].get("overall_risk")],
        ["Total Metrics", report["executive_summary"].get("total_metrics")],
        ["Passed", report["executive_summary"].get("passed")],
        ["Failed", report["executive_summary"].get("failed")],
        ["Critical Failures", report["executive_summary"].get("critical_failures")],
        ["Recommendation", report["executive_summary"].get("recommendation")],
    ])
    _docx_table(
        document,
        "Evaluation Summary",
        [["Metric", "Category", "Score", "Threshold", "Status", "Severity"]]
        + [[item["metric"], item["category"], item["score_display"], item["threshold_display"], item["status"], item["severity"]] for item in report.get("evaluation_summary") or []],
    )
    _docx_table(document, "Executive Scorecard", [["Metric", "Score", "Status", "Reason"]] + [[m["name"], m["display"], "PASS" if m["passed"] else "FAIL", m["reason"]] for m in report["scorecard"]["metrics"]])
    _docx_table(document, "Capability Matrix", [["Capability", "Status", "Evidence"]] + [[r["capability"], r["status"], r["evidence"]] for r in report["capability_matrix"]])
    document.add_heading("Conversation Timeline", level=1)
    for turn in report["conversation_timeline"][:60]:
        evidence = turn.get("interaction_evidence") if isinstance(turn.get("interaction_evidence"), dict) else None
        if evidence:
            title, subtitle = _interaction_title(evidence)
            p = document.add_paragraph()
            p.add_run(f"Turn {turn['turn']} · EvalForge Action · {title}: ").bold = True
            p.add_run(subtitle)
            _docx_evidence_image(document, evidence)
            _docx_table(document, title, [["Field", "Value"]] + _interaction_detail_rows(evidence))
            continue
        p = document.add_paragraph()
        p.add_run(f"Turn {turn['turn']} · {turn['speaker']}: ").bold = True
        p.add_run(turn["message"])
        _docx_turn_components(document, turn)
    _docx_table(document, "AI Findings", [["Finding", "Severity", "Score", "Summary"]] + [[f["title"], f["severity"], f["score"], f["summary"]] for f in report["ai_findings"]])
    _docx_table(document, "Risk Analysis", [["Level", "Risk", "Impact"]] + [[r["level"], r["risk"], r["impact"]] for r in (report["risks"] or [{"level": "N/A", "risk": "N/A", "impact": "N/A"}])])
    _docx_table(document, "Backend Validation", _backend_validation_rows(report.get("backend_validation")))
    if report.get("workflow_validation"):
        _docx_table(
            document,
            "Workflow Validation",
            [["Step", "Expected", "Observed", "Status"]]
            + [[row.get("step"), row.get("expected"), row.get("observed"), row.get("status")] for row in report.get("workflow_validation")],
        )
    if report.get("rag_validation") and (report["rag_validation"].get("metrics") or report["rag_validation"].get("questions")):
        _docx_table(
            document,
            "RAG Validation",
            [["Metric", "Score", "Status", "Evidence"]]
            + [[m.get("name"), _percent(m.get("score")), m.get("status"), m.get("reason")] for m in report["rag_validation"].get("metrics") or []],
        )
    _docx_failure_analysis(document, report)
    _docx_table(
        document,
        "Overall Findings",
        [["Category", "Issues", "Metrics Evaluated"]]
        + [[item.get("category"), item.get("issues"), item.get("metrics")] for item in report.get("overall_findings") or []],
    )
    document.add_heading("Priority Recommendations", level=1)
    for item in report.get("priority_recommendations") or []:
        document.add_paragraph(f"Priority {item.get('priority')}: {item.get('title')} - {item.get('recommendation')}", style="List Bullet")
    _docx_table(document, "Artifact Summary", [["Artifact", "Count"]] + [[k.replace("_", " ").title(), str(_artifact_count(v))] for k, v in report["artifacts"].items()])
    _docx_table(document, "Appendix", _appendix_summary_rows(report.get("appendix")))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_failure_analysis(document: Any, report: dict[str, Any]) -> None:
    failures = [item for item in _list(report.get("failure_analysis")) if isinstance(item, dict)]
    if not failures:
        return
    document.add_heading("Failure Analysis", level=1)
    for index, failure in enumerate(failures, start=1):
        metric = failure.get("metric_name") or f"Failure {index}"
        document.add_heading(_value(metric), level=2)
        _docx_table(document, "Failure Summary", [
            ["Field", "Value"],
            ["Status", failure.get("status") or "FAILED"],
            ["Severity", failure.get("severity") or "N/A"],
            ["Category", failure.get("category") or "Uncategorized"],
            ["Source", failure.get("source") or "N/A"],
            ["Reason", failure.get("reason") or "N/A"],
            ["Expected", _component_value(failure.get("expected"))],
            ["Actual", _component_value(failure.get("actual"))],
            ["Root Cause", failure.get("root_cause") or "N/A"],
            ["Recommendation", failure.get("recommendation") or "N/A"],
        ])
        snippet_rows = _failure_conversation_rows(failure)
        if snippet_rows:
            _docx_table(document, "Conversation Evidence", snippet_rows)
        _docx_readable_evidence(document, failure.get("evidence"))
        details = failure.get("technical_details")
        if details not in (None, "", [], {}):
            _docx_readable_technical_details(document, details, failure)


def _docx_failure_conversation_rows(document: Any, rows: Any) -> None:
    rows = _list(rows)
    if len(rows) <= 1:
        return
    document.add_heading("Conversation Evidence", level=2)
    for row in rows[1:]:
        if not isinstance(row, list) or len(row) < 3:
            continue
        p = document.add_paragraph()
        label = _value(row[0])
        if "failure" in label.lower():
            label += " - FAILURE OCCURRED HERE"
        p.add_run(f"{label} · {_value(row[1])}: ").bold = True
        p.add_run(_value(row[2]))


def _docx_readable_evidence(document: Any, evidence: Any) -> None:
    items = _list(evidence)
    if not items:
        return
    document.add_heading("Supporting Evidence", level=2)
    for index, item in enumerate(items, start=1):
        if isinstance(item, dict) and str(item.get("type") or "").lower() == "conversation_card":
            _docx_conversation_card_evidence(document, item)
            continue
        facts = _readable_evidence_facts(item) if isinstance(item, dict) else [("Evidence", _component_value(item))]
        if facts:
            _docx_table(document, f"Supporting Evidence {index}", [["Field", "Value"]] + [[label, value] for label, value in facts])


def _docx_conversation_card_evidence(document: Any, item: dict[str, Any]) -> None:
    document.add_heading("Conversation Evidence", level=2)
    for title, key, failed in (
        ("Previous Turn", "previous_turn", False),
        ("Failed Turn", "failed_turn", True),
        ("Next Turn", "next_turn", False),
    ):
        p = document.add_paragraph()
        label = title + (" - FAILURE OCCURRED HERE" if failed else "")
        p.add_run(f"{label}: ").bold = True
        p.add_run(_component_value(item.get(key) or "N/A"))
    why = item.get("why_this_proves_failure")
    if why:
        p = document.add_paragraph()
        p.add_run("Why this proves the failure: ").bold = True
        p.add_run(_component_value(why))


def _docx_readable_technical_details(document: Any, details: Any, failure: dict[str, Any]) -> None:
    details = details if isinstance(details, dict) else {}
    rows = _readable_technical_rows(details, failure)
    if rows:
        _docx_table(document, "Technical Details", [["Field", "Value"]] + [[label, value] for label, value in rows])
    supporting = _supporting_evidence_rows_from_details(details)
    if supporting:
        _docx_table(document, "Supporting Evidence", supporting)


def _docx_scorecards(document: Any, report: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    values = [
        ("Overall Score", _percent(report["header"]["overall_score"])),
        ("Status", report["header"]["status"]),
        ("Passed Metrics", str(report["scorecard"]["passed_metrics"])),
        ("Failed Metrics", str(report["scorecard"]["failed_metrics"])),
    ]
    for index, (label, value) in enumerate(values):
        cell = table.rows[0].cells[index]
        p = cell.paragraphs[0]
        p.add_run(label + "\n").bold = True
        p.add_run(value)
        
    categories = {}
    for metric in report["scorecard"]["metrics"]:
        cat = metric.get("category") or "Conversation"
        categories.setdefault(cat, []).append(metric)
        
    for cat, metrics in categories.items():
        document.add_heading(f"{str(cat).title()} Metrics", level=2)
        rows = [["Metric", "Score", "Status", "Reason"]]
        for m in metrics:
            rows.append([str(m["name"]), str(m["display"]), str(m["status"]), str(m["reason"])])
        _docx_table(document, f"{str(cat).title()} Metrics Table", rows)


def _docx_table(document: Any, title: str, rows: list[list[Any]]) -> None:
    document.add_heading(title, level=1)
    if not rows:
        document.add_paragraph("N/A")
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = _value(value)
            if r_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def _docx_turn_components(document: Any, turn: dict[str, Any]) -> None:
    components = [item for item in _list(turn.get("components")) if isinstance(item, dict)]
    grouped = _group_components(components)
    if grouped["options"]:
        _docx_table(
            document,
            "Available Options",
            [["Type", "Value"]] + [[_component_type_label(item), _option_value(item)] for item in grouped["options"]],
        )
    for component in grouped["other"]:
        type_name = str(component.get("type") or component.get("subtype") or "").lower()
        if type_name == "table":
            rows = _component_table_rows(component)
            if rows:
                _docx_table(document, _value(component.get("title") or "Table"), rows)
        elif type_name in {"form", "card"}:
            fields = component.get("fields") or component.get("items") or component.get("sections") or []
            _docx_table(
                document,
                _value(component.get("title") or component.get("label") or type_name.title()),
                [["Field", "Value"]] + _component_field_rows(fields),
            )
        else:
            for line in _turn_artifact_lines({"components": [component], "attachments": []}):
                document.add_paragraph(line, style="List Bullet")
    for attachment in _list(turn.get("attachments")):
        document.add_paragraph(_attachment_label(attachment), style="List Bullet")


def _pdf_evidence_image(story: list[Any], evidence: dict[str, Any]) -> None:
    local_path = _local_existing_path(_evidence_image_path(evidence))
    if not local_path:
        return
    try:
        from reportlab.lib.units import inch
        from reportlab.platypus import Image, Spacer

        story.append(Image(local_path, width=2.6 * inch, height=1.9 * inch, kind="proportional"))
        story.append(Spacer(1, 6))
    except Exception:
        return


def _docx_evidence_image(document: Any, evidence: dict[str, Any]) -> None:
    local_path = _local_existing_path(_evidence_image_path(evidence))
    if not local_path:
        return
    try:
        from docx.shared import Inches

        document.add_picture(local_path, width=Inches(2.6))
    except Exception:
        return


def _evidence_image_path(evidence: dict[str, Any]) -> str:
    mime_type = str(evidence.get("mime_type") or evidence.get("file_type") or "").lower()
    path = str(evidence.get("artifact_path") or evidence.get("preview_url") or evidence.get("thumbnail_url") or "")
    if not path:
        return ""
    if mime_type.startswith("image/") or _is_image_attachment({"path": path, "mime_type": mime_type}):
        return path
    return ""


def _local_existing_path(path: str) -> str:
    if not path:
        return ""
    if re.match(r"^(https?:|data:|blob:)", path, flags=re.I):
        return ""
    normalized = os.path.normpath(path)
    return normalized if os.path.exists(normalized) else ""


def _rgb(hex_color: str) -> Any:
    from docx.shared import RGBColor
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _collect_components(conversation: list[Any]) -> list[dict[str, Any]]:
    components: list[dict[str, Any]] = []
    for turn in conversation:
        if not isinstance(turn, dict):
            continue
        components.extend(_turn_components(turn))
    return components


def _failure_analysis(value: Any) -> list[dict[str, Any]]:
    failures: list[dict[str, Any]] = []
    for item in _list(value):
        if isinstance(item, dict):
            failures.append(item)
        elif hasattr(item, "model_dump"):
            dumped = item.model_dump()
            if isinstance(dumped, dict):
                failures.append(dumped)
    return failures


def _failure_conversation_rows(failure: dict[str, Any]) -> list[list[Any]]:
    snippet = [item for item in _list(failure.get("conversation_snippet")) if isinstance(item, dict)]
    if not snippet:
        return []
    rows: list[list[Any]] = [["Turn", "Speaker", "Message"]]
    failure_turn = failure.get("conversation_turn")
    for item in snippet:
        turn = item.get("turn") or item.get("conversation_turn") or "N/A"
        label = _value(item.get("label") or "")
        is_failure = "failure" in label.lower() or (failure_turn is not None and str(turn) == str(failure_turn))
        turn_label = f"Turn {turn}"
        if is_failure:
            turn_label += " - FAILURE OCCURRED HERE"
        speaker = item.get("speaker") or item.get("role") or "N/A"
        message = item.get("message") or item.get("content") or item.get("text") or item
        rows.append([turn_label, speaker, _component_value(message)])
    return rows


def _failure_evidence_rows(value: Any) -> list[list[Any]]:
    evidence = _list(value)
    if not evidence:
        return []
    rows: list[list[Any]] = [["Evidence", "Details"]]
    for index, item in enumerate(evidence, start=1):
        if isinstance(item, dict):
            kind = item.get("type") or item.get("source") or item.get("kind") or f"Evidence {index}"
            rows.append([kind, _component_value(item)])
        else:
            rows.append([f"Evidence {index}", _component_value(item)])
    return rows


def _turn_components(turn: Any) -> list[dict[str, Any]]:
    if not isinstance(turn, dict):
        return []
    metadata = turn.get("metadata") if isinstance(turn.get("metadata"), dict) else {}
    raw = turn.get("components") or metadata.get("components") or metadata.get("rich_components") or []
    return [item for item in _list(raw) if isinstance(item, dict)]


def _turn_attachments(turn: Any) -> list[Any]:
    if not isinstance(turn, dict):
        return []
    metadata = turn.get("metadata") if isinstance(turn.get("metadata"), dict) else {}
    return _list(turn.get("attachments") or metadata.get("attachments"))


def _interaction_evidence(turn: Any) -> dict[str, Any] | None:
    if not isinstance(turn, dict):
        return None
    metadata = turn.get("metadata") if isinstance(turn.get("metadata"), dict) else {}
    evidence = metadata.get("interaction_evidence")
    return evidence if isinstance(evidence, dict) else None


def _by_type(components: list[dict[str, Any]], types: set[str]) -> list[dict[str, Any]]:
    return [c for c in components if str(c.get("type") or c.get("subtype") or "").lower() in types]


def _attachments(conversation: list[dict[str, Any]], components: list[dict[str, Any]]) -> list[Any]:
    result: list[Any] = []
    for turn in conversation:
        if isinstance(turn, dict):
            result.extend(_turn_attachments(turn))
    result.extend(_by_type(components, {"attachment", "file", "image", "video"}))
    return result


def _otp(conversation: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for turn in conversation:
        if not isinstance(turn, dict):
            continue
        metadata = turn.get("metadata") if isinstance(turn.get("metadata"), dict) else {}
        evidence = metadata.get("interaction_evidence") if isinstance(metadata.get("interaction_evidence"), dict) else {}
        if str(evidence.get("interaction_type") or "").lower() in {"otp", "otp_entry"}:
            result.append({"speaker": turn.get("speaker"), "message": "***MASKED***", "status": evidence.get("status")})
        elif "otp" in str(turn.get("message") or "").lower() or metadata.get("action") == "otp":
            result.append({"speaker": turn.get("speaker"), "message": "***MASKED***"})
    return result


def _interaction_summary(conversation: list[dict[str, Any]]) -> dict[str, int]:
    summary = {
        "forms_detected": len(_by_type(_collect_components(conversation), {"form"})),
        "forms_submitted": 0,
        "images_uploaded": 0,
        "camera_captures": 0,
        "documents_uploaded": 0,
        "videos_uploaded": 0,
        "otp_verifications": 0,
        "buttons_clicked": 0,
        "dropdown_selections": 0,
        "radio_selections": 0,
        "checkbox_selections": 0,
        "date_selections": 0,
        "time_selections": 0,
        "human_handoffs": 0,
        "backend_validations": 0,
        "ui_validations": 0,
    }
    mapping = {
        "form_fill": "forms_submitted",
        "form": "forms_submitted",
        "image_upload": "images_uploaded",
        "camera_capture": "camera_captures",
        "file_upload": "documents_uploaded",
        "document_upload": "documents_uploaded",
        "video_upload": "videos_uploaded",
        "otp": "otp_verifications",
        "otp_entry": "otp_verifications",
        "button_click": "buttons_clicked",
        "dropdown_selection": "dropdown_selections",
        "radio_selection": "radio_selections",
        "checkbox_selection": "checkbox_selections",
        "multi_select": "checkbox_selections",
        "date_picker": "date_selections",
        "time_picker": "time_selections",
        "human_handoff": "human_handoffs",
        "backend_validation": "backend_validations",
        "ui_validation": "ui_validations",
    }
    for turn in conversation:
        evidence = _interaction_evidence(turn)
        if not evidence:
            continue
        key = mapping.get(str(evidence.get("interaction_type") or "").lower())
        if key:
            summary[key] += 1
    return summary


def _interaction_title(evidence: dict[str, Any]) -> tuple[str, str]:
    type_name = str(evidence.get("interaction_type") or "").lower()
    selected = evidence.get("selected_value")
    if type_name in {"form_fill", "form"}:
        return "Form Submitted", _value(evidence.get("action") or evidence.get("requested_by_agent"))
    if type_name == "camera_capture":
        return "Live Camera", "Captured successfully"
    if type_name == "image_upload":
        return "Image Uploaded", _value(evidence.get("filename") or evidence.get("source") or selected)
    if type_name in {"file_upload", "document_upload", "attachment_upload"}:
        return "Document Uploaded", _value(evidence.get("filename") or selected)
    if type_name == "video_upload":
        return "Video Uploaded", _value(evidence.get("filename") or selected)
    if type_name in {"otp", "otp_entry"}:
        return "OTP Verified", "******"
    if type_name == "button_click":
        return "Button Clicked", _value(selected)
    if type_name == "dropdown_selection":
        return "Selected Option", _value(selected)
    if type_name == "radio_selection":
        return "Selected", _value(selected)
    if type_name in {"checkbox_selection", "multi_select"}:
        return "Selected", _component_value(selected)
    if type_name == "date_picker":
        return "Date Selected", _value(selected)
    if type_name == "time_picker":
        return "Time Selected", _value(selected)
    if type_name == "search":
        return "Search Performed", _value(selected)
    if type_name == "backend_validation":
        return "Backend Validation", _value(evidence.get("action"))
    if type_name == "ui_validation":
        subtitle = evidence.get("form_name") or evidence.get("error") or "Browser UI verified"
        return "UI Validation", _value(subtitle)
    return _value(evidence.get("action") or type_name or "Interaction"), _component_value(selected)


def _interaction_detail_rows(evidence: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    type_name = str(evidence.get("interaction_type") or "").lower()
    selected = evidence.get("selected_value")
    if type_name == "form_fill" and isinstance(selected, dict):
        for key, value in selected.items():
            rows.append([_value(key), _mask_sensitive(key, value)])
    elif selected not in (None, "", [], {}):
        rows.append(["Selected Value", "******" if "otp" in type_name else _component_value(selected)])
    for label, value in [
        ("Action", evidence.get("action")),
        ("Status", evidence.get("status") or ("FAILED" if evidence.get("success") is False else "SUCCESS")),
        ("Requested By Agent", evidence.get("requested_by_agent")),
        ("Source", evidence.get("source")),
        ("Filename", evidence.get("filename")),
        ("File Type", evidence.get("mime_type") or evidence.get("file_type")),
        ("File Size", evidence.get("size_bytes") or evidence.get("size")),
        ("Fields Completed", evidence.get("fields_completed")),
        ("Required Fields", evidence.get("required_fields")),
        ("Optional Fields", evidence.get("optional_fields")),
        ("Form Detected", _yes_no(evidence.get("form_detected")) if "form_detected" in evidence else None),
        ("Form Name", evidence.get("form_name")),
        ("Fields Detected", evidence.get("fields_detected")),
        ("Submit Button", _yes_no(evidence.get("submit_button_detected")) if "submit_button_detected" in evidence else None),
        ("Upload Control", _yes_no(evidence.get("upload_control_detected")) if "upload_control_detected" in evidence else None),
        ("Accept", evidence.get("accept")),
        ("Camera Available", _yes_no(evidence.get("camera_available")) if "camera_available" in evidence else None),
        ("OTP Field", _yes_no(evidence.get("otp_field_detected")) if "otp_field_detected" in evidence else None),
        ("Date Field", _yes_no(evidence.get("date_field_detected")) if "date_field_detected" in evidence else None),
        ("Time Field", _yes_no(evidence.get("time_field_detected")) if "time_field_detected" in evidence else None),
        ("Search Field", _yes_no(evidence.get("search_field_detected")) if "search_field_detected" in evidence else None),
        ("Buttons Detected", evidence.get("buttons_detected")),
        ("Retakes", evidence.get("retakes")),
        ("Duration", f"{int(evidence.get('duration_ms')) / 1000:g} sec" if evidence.get("duration_ms") else None),
        ("Turn", evidence.get("turn_number")),
        ("Error", evidence.get("error") or evidence.get("reason")),
    ]:
        if value not in (None, "", [], {}):
            rows.append([label, _component_value(value)])
    return rows or [["Status", "N/A"]]


def _mask_sensitive(name: Any, value: Any) -> str:
    key = str(name or "").lower()
    if "otp" in key or "password" in key or "passcode" in key:
        return "******"
    return _component_value(value)


def _sanitize_report_json(value: Any, key: str = "") -> Any:
    lowered = str(key or "").lower()
    if any(token in lowered for token in ("otp", "password", "passcode", "secret", "token")):
        return "******" if value not in (None, "", [], {}) else value
    if isinstance(value, dict):
        is_otp_context = str(value.get("interaction_type") or value.get("action") or "").lower().find("otp") >= 0
        sanitized = {}
        for k, v in value.items():
            if is_otp_context and str(k).lower() in {"selected_value", "value", "input", "code"}:
                sanitized[k] = "******"
            else:
                sanitized[k] = _sanitize_report_json(v, k)
        return sanitized
    if isinstance(value, list):
        return [_sanitize_report_json(item, key) for item in value]
    if isinstance(value, str) and "otp" in lowered:
        return "******"
    return value


def _status(payload: dict[str, Any]) -> str:
    if payload.get("status"):
        return str(payload["status"])
    if payload.get("success") is False:
        return "ERROR"
    if payload.get("passed") is True:
        return "PASS"
    if payload.get("passed") is False:
        return "FAIL"
    return "N/A"


def _normalize_metric(metric: dict[str, Any], default_threshold: float | None = None) -> dict[str, Any]:
    normalized = dict(metric)
    score = _float(metric.get("score"))
    if score is None:
        score = _float(metric.get("value") or metric.get("metric_score") or metric.get("actual_score"))
    threshold = _float(metric.get("threshold") or metric.get("pass_threshold"))
    if threshold is None:
        threshold = default_threshold
    passed = None
    if score is not None and threshold is not None:
        passed = _score_for_compare(score) >= _score_for_compare(threshold)
    if passed is None:
        passed = metric.get("passed")
    if passed is None:
        passed = metric.get("success")
    if passed is None:
        status = str(metric.get("status") or metric.get("result") or metric.get("outcome") or "").strip().lower()
        if status in {"pass", "passed", "success", "successful", "ok"}:
            passed = True
        elif status in {"fail", "failed", "failure", "error", "not_passed"}:
            passed = False
    evaluation_status = str(metric.get("evaluation_status") or "").lower()
    if evaluation_status in {"not_evaluated", "error"}:
        passed = None
        score = None
    normalized["score"] = score
    normalized["passed"] = passed
    normalized["status"] = "NOT EVALUATED" if evaluation_status == "not_evaluated" else ("ERROR" if evaluation_status == "error" else ("PASS" if passed is True else ("FAIL" if passed is False else _value(metric.get("status")))))
    normalized["threshold"] = threshold
    normalized["name"] = metric.get("name") or metric.get("metric") or metric.get("label") or "Metric"
    normalized["category"] = metric.get("category") or _metric_category_from_name(normalized["name"])
    normalized["severity"] = metric.get("severity") or _metric_severity(normalized)
    normalized["reason"] = metric.get("reason") or metric.get("explanation") or metric.get("details") or ""
    return normalized


def _metric_category_from_name(name: Any) -> str:
    text = str(name or "").lower()
    if "intent" in text:
        return "Intent"
    if "persona" in text or "tone" in text:
        return "Persona"
    if "context" in text or "memory" in text:
        return "Context"
    if "relevancy" in text or "response" in text or "quality" in text:
        return "Response Validation"
    if "complete" in text or "goal" in text:
        return "Completeness"
    if "workflow" in text or "form" in text or "upload" in text:
        return "Workflow"
    if "rag" in text or "faithfulness" in text or "hallucination" in text:
        return "RAG"
    if "backend" in text:
        return "Backend Validation"
    return "Conversation"


def _metric_severity(metric: dict[str, Any]) -> str:
    if metric.get("passed") is True:
        return "None"
    score = _float(metric.get("score"))
    if score is None:
        return "Medium"
    score = _score_for_compare(score)
    if score < 25:
        return "Critical"
    if score < 50:
        return "High"
    if score < 70:
        return "Medium"
    return "Low"


def _metric_threshold(payload: dict[str, Any]) -> float | None:
    candidates = [
        payload.get("pass_threshold"),
        payload.get("threshold"),
        (payload.get("scenario") or {}).get("pass_threshold") if isinstance(payload.get("scenario"), dict) else None,
    ]
    for candidate in candidates:
        threshold = _float(candidate)
        if threshold is not None:
            return threshold
    return 0.7


def _turn_artifact_lines(turn: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    for component in _list(turn.get("components")):
        if not isinstance(component, dict):
            continue
        type_name = str(component.get("type") or component.get("subtype") or "component").lower()
        if type_name in {"button", "quick_reply", "option", "suggested_prompt", "suggested-prompt", "link"}:
            lines.append(f"{_component_type_label(component)}: {_option_value(component)}")
        elif type_name == "table":
            table_data = component.get("data") if isinstance(component.get("data"), dict) else component
            lines.append(f"Table: {len(_list(table_data.get('rows')))} row(s)")
        elif type_name in {"form", "card"}:
            lines.append(f"{type_name.title()}: {_value(component.get('title') or component.get('label'))}")
        elif type_name in {"image", "video", "attachment", "file"}:
            lines.append(f"{type_name.title()}: {_value(component.get('name') or component.get('filename') or component.get('url') or component.get('path'))}")
        else:
            lines.append(f"Component: {type_name}")
    for attachment in _list(turn.get("attachments")):
        lines.append(f"Attachment: {_value(attachment.get('name') if isinstance(attachment, dict) else attachment)}")
    return lines


def _column_label(column: Any) -> str:
    if isinstance(column, dict):
        return _value(column.get("label") or column.get("name") or column.get("key") or column.get("id"))
    return _value(column)


def _component_table_rows(component: dict[str, Any]) -> list[list[str]]:
    table_data = component.get("data") if isinstance(component.get("data"), dict) else component
    headers = _list(table_data.get("headers") or table_data.get("columns"))
    rows = _list(table_data.get("rows"))
    if not headers and rows and isinstance(rows[0], dict):
        headers = list(rows[0].keys())
    if not headers and rows and isinstance(rows[0], list):
        headers = [f"Column {index + 1}" for index in range(len(rows[0]))]
    if not headers:
        return []
    labels = [_column_label(header) for header in headers]
    keys = [
        header.get("key") or header.get("id") or header.get("name") or header.get("label")
        if isinstance(header, dict)
        else header
        for header in headers
    ]
    normalized_rows: list[list[str]] = []
    for row in rows:
        if isinstance(row, dict):
            normalized_rows.append([_component_value(row.get(key)) for key in keys])
        elif isinstance(row, list):
            normalized_rows.append([_component_value(cell) for cell in row])
        else:
            normalized_rows.append([_component_value(row)])
    return [labels] + normalized_rows


def _component_field_rows(fields: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for field in _list(fields):
        if isinstance(field, dict):
            rows.append([
                _value(field.get("label") or field.get("name") or field.get("id") or "Field"),
                _component_value(field.get("value") or field.get("text") or field.get("type") or ""),
            ])
        else:
            rows.append([_value(field), ""])
    return rows


def _component_value(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)
    return _value(value)


def _artifact_count(value: Any) -> int:
    if isinstance(value, int):
        return value
    if isinstance(value, list):
        return len(value)
    if isinstance(value, dict):
        return sum(item for item in value.values() if isinstance(item, int))
    return 0


def _group_components(items: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    options: list[dict[str, Any]] = []
    other: list[dict[str, Any]] = []
    for item in items:
        type_name = str(item.get("type") or item.get("subtype") or "").lower()
        if type_name in {"button", "quick_reply", "option", "suggested_prompt", "suggested-prompt"}:
            options.append(item)
        else:
            other.append(item)
    return {"options": options, "other": other}


def _component_type_label(component: dict[str, Any]) -> str:
    type_name = str(component.get("type") or component.get("subtype") or "Option").replace("_", " ").replace("-", " ").strip()
    return type_name.title()


def _option_value(component: dict[str, Any]) -> str:
    label = component.get("label") or component.get("text") or component.get("message") or component.get("value") or component.get("title") or "Option"
    prefix = component.get("field") or component.get("group") or component.get("context")
    return f"{prefix}: {label}" if prefix else _value(label)


def _attachment_label(item: Any) -> str:
    component = item if isinstance(item, dict) else {"path": item, "name": item}
    label = component.get("name") or component.get("filename") or component.get("file_name") or component.get("path") or component.get("url") or "Attachment"
    return _value(label)


def _is_image_attachment(item: Any) -> bool:
    component = item if isinstance(item, dict) else {"path": item, "name": item}
    mime = str(component.get("mime_type") or component.get("content_type") or component.get("type") or "").lower()
    if mime.startswith("image/"):
        return True
    candidate = str(
        component.get("name")
        or component.get("filename")
        or component.get("file_name")
        or component.get("preview_url")
        or component.get("thumbnail_url")
        or component.get("url")
        or component.get("src")
        or component.get("path")
        or component.get("blob_url")
        or ""
    ).lower().split("?", 1)[0].split("#", 1)[0]
    return bool(re.search(r"\.(png|jpe?g|gif|webp|bmp|svg)$", candidate))


def _looks_like_otp(component: dict[str, Any]) -> bool:
    haystack = " ".join(str(component.get(key, "")) for key in ("type", "subtype", "label", "title", "action", "name")).lower()
    return "otp" in haystack or "verification" in haystack


def _yes_no(value: Any) -> str:
    if value is True:
        return "Yes"
    if value is False:
        return "No"
    return "N/A"


def _score_for_compare(value: float) -> float:
    return value / 100 if value > 1 else value


def _bar_width(value: Any) -> int:
    number = _float(value)
    if number is None:
        return 0
    percent = number if number > 1 else number * 100
    return max(0, min(100, int(percent)))


def _safe_url(value: Any) -> str:
    if not value:
        return ""
    url = str(value).strip()
    if re.match(r"^(https?:|data:image/|blob:)", url, flags=re.I):
        return url
    if re.match(r"^[./A-Za-z0-9_\-/%]+$", url) and not re.match(r"^[a-z]:", url, flags=re.I):
        return url
    return ""


def _value(value: Any) -> str:
    if value is None or value == "":
        return "N/A"
    return str(value)


def _float(value: Any) -> float | None:
    try:
        if value is None:
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _percent(value: Any) -> str:
    number = _float(value)
    if number is None:
        return "N/A"
    percent = number if number > 1 else number * 100
    return f"{percent:.1f}%"


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _now() -> str:
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def e(value: Any) -> str:
    return html.escape(_value(value), quote=True)


def _xml(value: Any) -> str:
    return html.escape(_value(value), quote=False)


def _wrap(text: Any, width: int) -> list[str]:
    words = re.sub(r"\s+", " ", _value(text)).split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        if len(current) + len(word) + 1 > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    if current:
        lines.append(current)
    return lines or [""]


def _pdf_escape(text: str) -> str:
    return str(text).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
