"""HTML, PDF, DOCX and CSV exporters for consolidated management reports."""

from __future__ import annotations

import csv
import html
import io
from typing import Any

from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer

from app.management.consolidated_models import ConsolidatedReport


_OUTCOME_EMPTY = "No PASS, FAIL, or ERROR evaluation results are available for the selected report scope."


def _outcome_rows(d):
    return [[item["outcome"], item["count"], f"{item['percentage']:.1f}%"] for item in d["outcome_distribution"]["slices"]]


def _outcome_notes(d):
    distribution = d["outcome_distribution"]
    notes = [f"Total represented: {distribution['represented_total']}"]
    if distribution["unknown_count"]:
        notes.append(f"Unknown / unclassified: {distribution['unknown_count']} (excluded from the pie distribution).")
    return notes


def _outcome_html(d):
    empty = f"<p>{_OUTCOME_EMPTY}</p>" if not d["outcome_distribution"]["represented_total"] else ""
    return '<h3>Evaluation Outcome Distribution</h3>' + empty + _table(
        ["Outcome", "Count", "Percentage"], [[_e(cell) for cell in row] for row in _outcome_rows(d)]
    ) + ''.join(f"<p>{_e(note)}</p>" for note in _outcome_notes(d))


def _workflow_content(d, cases=False):
    if d.get("workflows") and not cases:
        yield "body", "Authored coverage uses the current catalog. Execution coverage and outcomes use only selected evaluation runs.", None
    for workflow in d.get("workflows", []):
        yield "heading", f"Workflow: {workflow['workflow_name']}", None
        if not cases:
            yield "body", f"Authored coverage: {workflow['authored_total']} / {workflow['required_total']} | {workflow['coverage_status']}", None
            if workflow["workflow_id"]:
                for kind, title in (("positive", "Positive"), ("negative", "Negative"), ("edge", "Edge / Adverse")):
                    yield "body", f"{title}: {workflow['authored'][kind]} / {workflow['required'][kind]} | Missing: {workflow['missing'][kind]}", None
            yield "body", f"Distinct cases in selected runs: attempted {workflow['attempted']}; evaluated at least once {workflow['evaluated']}; latest technical ERROR {workflow['technical_errors']}.", None
            yield "body", "Selected evaluation runs: " + "; ".join(f"{key.upper()} {value}" for key, value in workflow["outcome_counts"].items()), None
        else:
            for case in workflow["cases"]:
                kind = (case["test_case_type"] or "Legacy / Unclassified").replace("edge", "Edge / Adverse")
                yield "body", f"Test Case: {case['test_case_id']} | Type: {kind} | Version: {case['test_case_version'] or 'unrecorded'} | Outcome: {(case['outcome'] or 'NO SELECTED RUN').upper()} | Score: {_pct(case['score'])}", None
                if case["evaluation_id"]:
                    yield "body", f"Evaluation: {case['evaluation_id']} | Workflow snapshot: {case['workflow_name']}", None
                    yield "reference", case["report_url"], case["evaluation_id"]
    if not cases:
        for issue in d.get("catalog_issues", []):
            yield "body", "Catalog warning: " + _e(issue.get("code")) + " | " + str(issue.get("id") or issue.get("file") or "") + " | " + ", ".join(issue.get("files", [])), None


def _workflow_html(d, cases=False):
    parts = []
    for kind, text, evaluation_id in _workflow_content(d, cases):
        if kind == "reference":
            parts.append(_link(text, "View Report"))
        else:
            tag = "h3" if kind == "heading" else "p"
            parts.append(f'<{tag} style="overflow-wrap:anywhere">{_e(text)}</{tag}>')
    return ''.join(parts)


def _workflow_pdf(d, styles, cases=False):
    return [_pdf_report_reference(eid, d, styles) if kind == "reference" else Paragraph(_e(text), styles["Subheading" if kind == "heading" else "BodyText"])
            for kind, text, eid in _workflow_content(d, cases)]


def _workflow_docx(doc, d, cases=False):
    for kind, text, eid in _workflow_content(d, cases):
        if kind == "heading":
            doc.add_heading(text, 2)
        else:
            doc.add_paragraph(text)


def render_consolidated_html(report: ConsolidatedReport) -> str:
    data = report.model_dump()
    sections = report.section_order
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{_e(report.agent_name)} — Consolidated Evaluation Report</title>
<style>
body{{font-family:Arial,sans-serif;color:#172033;background:#f3f6fb;margin:0}}main{{max-width:1120px;margin:auto;padding:36px}}header{{background:#12213c;color:white;padding:34px;border-radius:8px}}h1{{margin:0 0 8px}}section{{background:white;margin:20px 0;padding:24px;border:1px solid #dbe3ef;border-radius:8px;break-inside:avoid}}h2{{font-size:20px;color:#173f73;border-bottom:2px solid #2e73c5;padding-bottom:10px}}.grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}}.card{{border:1px solid #dbe3ef;padding:14px;border-radius:6px}}.label{{font-size:11px;text-transform:uppercase;color:#627188}}.value{{font-size:24px;font-weight:700;margin-top:5px}}table{{border-collapse:collapse;width:100%;font-size:12px}}th,td{{border:1px solid #dbe3ef;padding:9px;text-align:left;vertical-align:top}}th{{background:#eaf1fa;color:#173f73}}.pass{{color:#15803d}}.fail{{color:#b91c1c}}.error{{color:#b45309}}.muted{{color:#627188}}.warning{{border-left:4px solid #d97706;background:#fff7ed;padding:12px;margin:8px 0}}a{{color:#1d5fa7}}@media(max-width:760px){{main{{padding:16px}}.grid{{grid-template-columns:1fr 1fr}}}}
</style></head><body><main><header><div>EvalForge Management</div><h1>Consolidated Evaluation Report</h1><p>{_e(report.agent_name)} · {_e(report.agent_type)} · {_e(report.provider or report.adapter or 'N/A')}</p><p>Generated {_e(report.generated_at)}</p></header>
{_html_section(sections[0], _executive_html(data))}
{_html_section(sections[1], _workflow_html(data) + _coverage_html(data))}
{_html_section(sections[2], _metrics_html(data))}
{_html_section(sections[3], _workflow_html(data, True) + _scenarios_html(data))}
{_html_section(sections[4], _issues_html(data))}
{_html_section(sections[5], _recommendations_html(data))}
{_html_section(sections[6], _history_html(data))}
{_html_section(sections[7], _references_html(data))}
</main></body></html>"""


def render_consolidated_pdf(report: ConsolidatedReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Brand", parent=styles["Heading2"], fontSize=11, leading=14, textColor=colors.HexColor("#2E73C5"), spaceAfter=8))
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontSize=27, leading=30, textColor=colors.HexColor("#173F73"), alignment=TA_LEFT, spaceAfter=8))
    styles.add(ParagraphStyle(name="CoverMeta", parent=styles["BodyText"], fontSize=9, leading=13, textColor=colors.HexColor("#334155")))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontSize=15, leading=18, textColor=colors.HexColor("#173F73"), borderColor=colors.HexColor("#2E73C5"), borderWidth=0, borderPadding=(0, 0, 5, 0), spaceBefore=10, spaceAfter=9, keepWithNext=True))
    styles.add(ParagraphStyle(name="Subheading", parent=styles["Heading3"], fontSize=10, leading=13, textColor=colors.HexColor("#334155"), spaceBefore=6, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8, leading=11, textColor=colors.HexColor("#334155")))
    styles.add(ParagraphStyle(name="Tiny", parent=styles["BodyText"], fontSize=7, leading=9, textColor=colors.HexColor("#475569")))
    styles.add(ParagraphStyle(name="Link", parent=styles["Tiny"], textColor=colors.HexColor("#1D5FA7")))
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=15*mm, rightMargin=15*mm, topMargin=16*mm, bottomMargin=17*mm, title="EvalForge Consolidated Evaluation Report", author="EvalForge")
    data = report.model_dump()
    period = data["executive_summary"].get("selected_period") or {}
    story: list[Any] = [
        Paragraph("EVALFORGE", styles["Brand"]),
        Paragraph("CONSOLIDATED<br/>EVALUATION REPORT", styles["CoverTitle"]),
        Spacer(1, 3*mm),
        _pdf_rule(Table, TableStyle, colors),
        Spacer(1, 4*mm),
        Paragraph(f"<b>Agent:</b> {_e(report.agent_name)} &nbsp;&nbsp; <b>Type:</b> {_e(report.agent_type)} &nbsp;&nbsp; <b>Adapter:</b> {_e(report.provider or report.adapter or 'N/A')}", styles["CoverMeta"]),
        Paragraph(f"<b>Evaluation Period:</b> {_e(period.get('from'))} to {_e(period.get('to'))} &nbsp;&nbsp; <b>Generated:</b> {_e(report.generated_at)}", styles["CoverMeta"]),
        Spacer(1, 7*mm),
    ]
    section_bodies = [
        _pdf_executive(data, styles, Table, TableStyle, colors),
        _pdf_coverage(data, styles, Table, TableStyle, colors),
        _pdf_metrics(data, styles, Table, TableStyle, colors),
        _pdf_scenarios(data, styles, Table, TableStyle, colors),
        _pdf_issues(data, styles),
        _pdf_recommendations(data, styles),
        _pdf_history(data, styles),
        _pdf_references(data, styles, Table, TableStyle, colors),
    ]
    for title, body in zip(report.section_order, section_bodies):
        story.append(CondPageBreak(28*mm))
        story.append(Paragraph(_e(title), styles["SectionTitle"]))
        story.append(_pdf_rule(Table, TableStyle, colors))
        story.append(Spacer(1, 3))
        story.extend(body)
        story.append(Spacer(1, 10))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#627188"))
        canvas.drawString(15*mm, 8*mm, f"EvalForge • Consolidated Evaluation Report • {report.agent_name}")
        canvas.drawRightString(landscape(A4)[0] - 15*mm, 8*mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_consolidated_docx(report: ConsolidatedReport) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    title = document.add_heading("Consolidated Evaluation Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(23, 63, 115)
    subtitle = document.add_paragraph(f"{report.agent_name}\n{report.agent_type} · {report.provider or report.adapter or 'N/A'}\nGenerated {report.generated_at}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    data = report.model_dump()
    builders = [_docx_executive, _docx_coverage, _docx_metrics, _docx_scenarios, _docx_issues, _docx_recommendations, _docx_history, _docx_references]
    for title_text, builder in zip(report.section_order, builders):
        document.add_heading(title_text, level=1)
        builder(document, data)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_consolidated_csv(report: ConsolidatedReport) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=["section", "record_type", "name", "value", "status", "score", "scenario_id", "evaluation_id", "details"])
    writer.writeheader()
    data = report.model_dump()
    for index, title in enumerate(report.section_order, 1):
        writer.writerow({"section": title, "record_type": "section", "name": title})
        if index == 1:
            for key, value in data["executive_summary"].items():
                writer.writerow({"section": title, "record_type": "summary", "name": key, "value": _csv_value(value)})
            for outcome, count, percentage in _outcome_rows(data):
                writer.writerow({"section": title, "record_type": "outcome_summary", "name": outcome, "value": count, "details": percentage})
        elif index == 2:
            for workflow in data.get("workflows", []):
                writer.writerow({"section": title, "record_type": "workflow_coverage", "name": workflow["workflow_name"], "value": workflow["authored_total"], "status": workflow["coverage_status"], "details": "; ".join(f"{k}={workflow['authored'][k]}/{v}" for k,v in workflow["required"].items())})
            for key, value in data["evaluation_coverage"].items():
                if key not in {"metric_coverage", "scenario_coverage"}:
                    writer.writerow({"section": title, "record_type": "coverage", "name": key, "value": _csv_value(value)})
        elif index == 3:
            for item in data["consolidated_metrics"]:
                writer.writerow({"section": title, "record_type": "metric", "name": item["name"], "score": item["average_score"], "details": f"evaluated={item['evaluated_count']}; pass={item['pass_count']}; fail={item['fail_count']}; not_evaluated={item['not_evaluated_count']}; error={item['error_count']}"})
        elif index == 4:
            for workflow in data.get("workflows", []):
                for case in workflow["cases"]:
                    writer.writerow({"section": title, "record_type": "workflow_case", "name": workflow["workflow_name"], "scenario_id": case["test_case_id"], "evaluation_id": case["evaluation_id"], "status": case["outcome"] or "not_run", "score": case["score"], "details": f"type={case['test_case_type'] or 'legacy'}; version={case['test_case_version'] or 'unrecorded'}"})
            for item in data["scenario_results"]:
                writer.writerow({"section": title, "record_type": "scenario", "scenario_id": item["scenario_id"], "evaluation_id": item["latest_valid_evaluation_id"], "score": item["latest_valid_score"], "status": item["latest_valid_outcome"], "details": f"runs={item['total_runs']}; valid={item['valid_runs']}; errors={item['technical_errors']}"})
        elif index == 5:
            for scenario in data["scenario_failures"]:
                for run in scenario["evaluations"]:
                    writer.writerow({"section": title, "record_type": "scenario_failure", "scenario_id": scenario["scenario_id"], "evaluation_id": run["evaluation_id"], "score": run["overall_score"], "status": "fail", "details": run["summary"]})
                    for metric in run["failed_metrics"]:
                        details = [metric["reason"]] + metric["failure_summaries"] + metric["evidence"] + metric["recommendations"]
                        writer.writerow({"section": title, "record_type": "failed_metric", "name": metric["name"], "scenario_id": scenario["scenario_id"], "evaluation_id": run["evaluation_id"], "score": metric["score"], "status": metric["status"], "details": "; ".join(filter(None, details))})
            for item in data["critical_failures"]:
                if item["technical"]:
                    writer.writerow({"section": title, "record_type": "issue", "name": item["category"], "status": "technical_error", "scenario_id": item["scenario_id"], "evaluation_id": item["evaluation_id"], "details": item["evidence"]})
        elif index == 6:
            for item in data["recommendations"]:
                writer.writerow({"section": title, "record_type": "recommendation", "name": item["trigger"], "value": item["recommendation"], "status": item["priority"], "details": ", ".join(item["evaluation_ids"])})
        elif index == 7:
            writer.writerow({"section": title, "record_type": "history", "status": data["historical_comparison"]["status"], "value": data["historical_comparison"]["message"], "details": "; ".join(data["historical_comparison"]["reasons"])})
        else:
            for item in data["individual_reports"]:
                writer.writerow({"section": title, "record_type": "evaluation", "scenario_id": item["scenario_id"], "evaluation_id": item["evaluation_id"], "status": item["outcome"], "score": item["score"], "details": item["report_url"]})
    return output.getvalue().encode("utf-8-sig")


def _html_section(title, content): return f"<section><h2>{_e(title)}</h2>{content}</section>"
def _pct(value): return "—" if value is None else f"{round(float(value)*100, 1)}%"
def _e(value): return html.escape(str(value if value not in (None, "") else "—"))
def _cards(items): return '<div class="grid">' + ''.join(f'<div class="card"><div class="label">{_e(k)}</div><div class="value">{_e(v)}</div></div>' for k,v in items) + '</div>'
def _table(headers, rows): return '<table><thead><tr>' + ''.join(f'<th>{_e(h)}</th>' for h in headers) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join(f'<td>{cell}</td>' for cell in row) + '</tr>' for row in rows) + '</tbody></table>'
def _link(url, label="View Individual Report"): return f'<a href="{_e(url)}">{_e(label)}</a>' if url else "—"


def _executive_html(d):
    x=d["executive_summary"]
    period=x.get("selected_period") or {}; period_text=f"{period.get('from') or '—'} to {period.get('to') or '—'}"
    cards=[("Mean Selected Run Score",_pct(x["mean_selected_run_score"])),("Total Evaluation Runs",x["total_evaluation_runs"]),("Unique Scenarios",x["unique_scenarios"]),("Passed Runs",x["passed_runs"]),("Failed Runs",x["failed_runs"]),("Technical Errors",x["technical_errors"]),("Unknown Outcomes",x["unknown_outcomes"]),("Selected Evaluation Period",period_text)]
    highlights=[("Strongest evaluated metric",x.get("strongest_evaluated_metric")),("Weakest evaluated metric",x.get("weakest_evaluated_metric")),("Most frequent quality issue",x.get("most_frequent_quality_issue")),("Technical reliability",x.get("technical_reliability"))]
    return _cards(cards)+_outcome_html(d)+f'<h3>Performance Overview</h3><p>{_e(x["performance_overview"])}</p><h3>Key Highlights</h3><ul>'+''.join(f'<li><b>{_e(k)}:</b> {_e(v)}</li>' for k,v in highlights if v)+'</ul>'


def _coverage_html(d):
    x=d["evaluation_coverage"]
    cards=[("Unique Scenarios Tested",x["unique_scenarios_tested"]),("Total Evaluation Runs",x["total_evaluation_runs"]),("Valid Scored Runs",x["valid_scored_runs"]),("Technical Error Runs",x["technical_error_runs"]),("Unknown Outcome Runs",x["unknown_outcome_runs"])]
    scenario_rows=[[_e(i["scenario_id"]),_e(i["total_runs"]),_e(i["valid_runs"]),_e(i["technical_errors"]),_e(i["last_evaluated"])] for i in x["scenario_coverage"]]
    warning=''.join(f'<div class="warning">{_e(w)}</div>' for w in x["comparability_warnings"]) or '<p class="muted">No comparability warning was persisted.</p>'
    metric_rows=[[_e(i["name"]),_e(i["evaluated_count"]),_e(i["not_evaluated_count"]),_e(i["error_count"]),_e(i["coverage"])] for i in x["metric_coverage"]]
    return _cards(cards)+'<h3>Outcome Distribution</h3>'+_cards([(k.upper(),v) for k,v in x["outcome_distribution"].items()])+'<h3>Scenario Coverage</h3>'+_table(["Scenario","Total Runs","Valid Runs","Technical Errors","Last Evaluated"],scenario_rows)+'<h3>Metric Coverage</h3>'+_table(["Metric","Evaluated","NOT EVALUATED","ERROR","Coverage"],metric_rows)+'<h3>Comparability</h3>'+warning


def _metrics_html(d):
    rows=[[_e(i["name"]),_e(_pct(i["average_score"])),_e(i["evaluated_count"]),_e(i["pass_count"]),_e(i["fail_count"]),_e(i["not_evaluated_count"]),_e(i["error_count"]),_e(i["coverage"])] for i in d["consolidated_metrics"]]
    return _table(["Metric","Average","Evaluated","Pass","Fail","NOT EVALUATED","ERROR","Coverage"],rows)


def _scenarios_html(d):
    rows=[[_e(i["scenario_id"]),_e(i["total_runs"]),_e(i["valid_runs"]),_e(i["technical_errors"]),_e(i["latest_valid_evaluation_id"]),_e(_pct(i["latest_valid_score"])),_e((i["latest_valid_outcome"] or "—").upper()),_e(i["last_evaluated"]),_link(i["report_url"])] for i in d["scenario_results"]]
    return _table(["Scenario","Runs","Valid","Errors","Latest Valid Evaluation","Score","Outcome","Last Evaluated","Report"],rows) if rows else '<p>No evaluations match this selection.</p>'


def _issues_html(d):
    parts = []
    for kind, text, evaluation_id in _failure_content(d):
        if kind == "reference":
            parts.append(_link(text, "View Report"))
        else:
            tag = "h3" if kind == "heading" else "p"
            parts.append(f'<{tag} style="overflow-wrap:anywhere">{_e(text)}</{tag}>')
    return ''.join(parts)


def _failure_content(d):
    """Shared export presentation of the already normalized scenario failures."""
    if not d.get("scenario_failures"):
        yield "body", "No failed test scenarios were found for the selected report scope.", None
    for scenario in d.get("scenario_failures", []):
        yield "heading", f"Scenario: {scenario['scenario_id']}", None
        for run in scenario["evaluations"]:
            metadata = run.get("scenario_metadata") or {}
            yield "body", f"Workflow: {metadata.get('workflow_name') or 'Legacy / Unmapped'} | Type: {metadata.get('test_case_type') or 'Legacy / Unclassified'} | Test Case: {scenario['scenario_id']}", None
            yield "heading", f"FAIL | Overall Score: {_pct(run['overall_score'])} | Evaluation: {run['evaluation_id']}", None
            yield "body", f"Failure Summary: {run['summary']}", None
            for metric in run["failed_metrics"]:
                yield "heading", f"{metric['name']} | FAIL | Score: {_pct(metric['score'])}", None
                if metric["reason"]:
                    yield "body", f"Reason: {metric['reason']}", None
                for summary in metric["failure_summaries"]:
                    yield "body", f"Failure Summary: {summary}", None
                for evidence in metric["evidence"]:
                    yield "body", f"Supporting Evidence: {evidence}", None
                for recommendation in metric["recommendations"]:
                    yield "body", f"Recommended Action: {recommendation}", None
            yield "reference", run["report_url"], run["evaluation_id"]
    technical = [i for i in d["critical_failures"] if i["technical"]]
    if technical:
        yield "heading", "Technical Errors (not scenario FAIL)", None
    for issue in technical:
        yield "heading", f"ERROR | {issue['scenario_id']} | {issue['evaluation_id']}", None
        yield "body", issue["evidence"], None
        yield "reference", issue["report_url"], issue["evaluation_id"]


def _recommendations_html(d):
    rows=[[_e(i["priority"]),_e(i["recommendation"]),_e(i["trigger"]),_e(', '.join(i["affected_scenarios"])),_e(', '.join(i["evaluation_ids"]))] for i in d["recommendations"]]
    return _table(["Priority","Recommendation","Trigger","Affected Scenarios","Supporting Evaluations"],rows) if rows else '<p>No evidence-backed recommendations were generated.</p>'


def _history_html(d):
    x=d["historical_comparison"]
    if x.get("status") == "comparable":
        return _cards([("Previous Batch",x.get("previous_batch")),("Current Batch",x.get("current_batch")),("Previous Mean",_pct(x.get("previous_mean_score"))),("Current Mean",_pct(x.get("current_mean_score"))),("Mean Change",_pct(x.get("mean_score_change")))])+_table(["Metric","Previous","Current","Change"],[[_e(i["metric"]),_e(_pct(i["previous"])),_e(_pct(i["current"])),_e(_pct(i["change"]))] for i in x.get("metric_changes",[])])
    return f'<div class="warning"><b>{_e(x["message"])}</b><ul>'+''.join(f'<li>{_e(r)}</li>' for r in x["reasons"])+'</ul></div>'


def _references_html(d):
    rows=[[_e(i["scenario_id"]),_e(i["evaluation_id"]),_e(i["timestamp"]),_e(i["outcome"].upper()),_e(_pct(i["score"])),_e(i["technical_status"]),_link(i["report_url"])] for i in d["individual_reports"]]
    return _table(["Scenario","Evaluation ID","Timestamp","Outcome","Score","Technical Status","Report"],rows) if rows else '<p>No included evaluation runs.</p>'


def _pdf_table(headers, rows, styles, Table, TableStyle, colors, widths=None):
    def cell(value):
        return value if isinstance(value, Paragraph) else Paragraph(_e(value), styles["Small"])
    body=[[cell(value) for value in headers]]+[[cell(value) for value in row] for row in rows]
    table=Table(body,colWidths=widths,repeatRows=1,hAlign="LEFT",splitByRow=1)
    commands=[("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DDE9F7")),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#173F73")),("LINEBELOW",(0,0),(-1,0),0.8,colors.HexColor("#7A9BC2")),("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#CBD5E1")),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]
    for row_index in range(2, len(body), 2):
        commands.append(("BACKGROUND",(0,row_index),(-1,row_index),colors.HexColor("#F8FAFC")))
    table.setStyle(TableStyle(commands))
    return table


def _pdf_rule(Table, TableStyle, colors):
    rule = Table([[""]], colWidths=[255*mm], rowHeights=[1.2*mm])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#2E73C5")), ("BOX", (0, 0), (-1, -1), 0, colors.white)]))
    return rule


def _pdf_reference_map(data):
    return {item["evaluation_id"]: item.get("report_url") for item in data["individual_reports"] if _safe_offline_report_path(item.get("report_url"))}


def _safe_offline_report_path(value):
    path = str(value or "")
    return path if path.startswith("Individual-Reports/") and ".." not in path and "\\" not in path else None


def _pdf_report_reference(evaluation_id, data, styles):
    evaluation_id = str(evaluation_id or "")
    path = _pdf_reference_map(data).get(evaluation_id)
    visible_id = evaluation_id if len(evaluation_id) <= 28 else evaluation_id[:25] + "..."
    if not path:
        return Paragraph(_e(visible_id or "—"), styles["Tiny"])
    return Paragraph(f'<link href="{_e(path)}" color="#1D5FA7"><u>View Individual Report</u></link><br/><font color="#475569">{_e(path)}</font>', styles["Link"])


def _pdf_status_label(value):
    text = str(value or "N/A").upper().replace("_", " ")
    return text if text else "N/A"


def _pdf_pct(value):
    return "N/A" if value is None else _pct(value)


def _pdf_executive(d,s,T,TS,c):
    x=d["executive_summary"]
    cards=[["MEAN SCORE","TOTAL RUNS","SCENARIOS","PASSED","FAILED","ERRORS","UNKNOWN"],[_pdf_pct(x["mean_selected_run_score"]),str(x["total_evaluation_runs"]),str(x["unique_scenarios"]),f"PASS {x['passed_runs']}",f"FAIL {x['failed_runs']}",f"ERROR {x['technical_errors']}",f"UNKNOWN {x['unknown_outcomes']}"]]
    table=T([[Paragraph(f"<b>{_e(label)}</b><br/><font size='14'>{_e(value)}</font>",s["Small"]) for label,value in zip(cards[0],cards[1])]],colWidths=[36*mm]*7,hAlign="LEFT")
    table.setStyle(TS([("BACKGROUND",(0,0),(-1,-1),c.HexColor("#F1F5F9")),("BOX",(0,0),(-1,-1),0.5,c.HexColor("#CBD5E1")),("INNERGRID",(0,0),(-1,-1),0.35,c.HexColor("#CBD5E1")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    return [table, *_pdf_outcome_distribution(d,s,c), Paragraph("Performance Overview",s["Subheading"]),Paragraph(_e(x["performance_overview"]),s["BodyText"])]


def _pdf_outcome_distribution(d, styles, colors):
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, Rect, String

    distribution = d["outcome_distribution"]
    blocks = [Paragraph("Evaluation Outcome Distribution", styles["Subheading"])]
    if distribution["represented_total"]:
        drawing = Drawing(440, 136)
        pie = Pie()
        pie.x, pie.y, pie.width, pie.height = 4, 6, 120, 120
        pie.innerRadiusFraction = .62
        visible = [item for item in distribution["slices"] if item["count"] > 0]
        pie.data = [item["count"] for item in visible]
        pie.labels = []
        for index, item in enumerate(visible):
            pie.slices[index].fillColor = colors.HexColor(item["color"])
            pie.slices[index].strokeColor = None
        drawing.add(pie)
        drawing.add(String(64, 66, str(distribution["represented_total"]), textAnchor="middle", fontName="Helvetica-Bold", fontSize=18))
        drawing.add(String(64, 51, "represented", textAnchor="middle", fontSize=8))
        for index, item in enumerate(distribution["slices"]):
            y = 105 - index * 30
            drawing.add(Rect(162, y-2, 10, 10, fillColor=colors.HexColor(item["color"]), strokeColor=None))
            drawing.add(String(182, y, item["outcome"], fontName="Helvetica-Bold", fontSize=10))
            drawing.add(String(315, y, str(item["count"]), textAnchor="end", fontSize=10))
            drawing.add(String(400, y, f"{item['percentage']:.1f}%", textAnchor="end", fontSize=10))
        blocks.append(drawing)
    else:
        blocks.append(Paragraph(_OUTCOME_EMPTY, styles["BodyText"]))
        for outcome, count, percentage in _outcome_rows(d):
            blocks.append(Paragraph(f"{outcome}: {count} ({percentage})", styles["Small"]))
    blocks.extend(Paragraph(_e(note), styles["Small"]) for note in _outcome_notes(d))
    return blocks
def _pdf_coverage(d,s,T,TS,c):
    x=d["evaluation_coverage"]; rows=[[i["scenario_id"],i["total_runs"],i["valid_runs"],i["technical_errors"],i["last_evaluated"] or "—"] for i in x["scenario_coverage"]]
    items=_workflow_pdf(d,s)+[Paragraph(f"Unique scenarios: {x['unique_scenarios_tested']} · Total runs: {x['total_evaluation_runs']} · Valid scored: {x['valid_scored_runs']} · Technical errors: {x['technical_error_runs']} · Unknown: {x['unknown_outcome_runs']}",s["BodyText"])]
    if rows:items.append(_pdf_table(["Scenario","Runs","Valid","Errors","Last Evaluated"],rows,s,T,TS,c,[65*mm,20*mm,20*mm,20*mm,75*mm]))
    metric_rows=[[i["name"],i["evaluated_count"],i["not_evaluated_count"],i["error_count"],i["coverage"]] for i in x["metric_coverage"]]
    items.extend([Paragraph("Metric Coverage",s["Heading3"]),_pdf_table(["Metric","Evaluated","Not Evaluated","Error","Coverage"],metric_rows,s,T,TS,c,[75*mm,28*mm,34*mm,24*mm,30*mm])])
    items.extend(Paragraph(_e(w),s["Small"]) for w in x["comparability_warnings"]); return items
def _pdf_metrics(d,s,T,TS,c):
    rows=[[i["name"],_pdf_pct(i["average_score"]),i["evaluated_count"],f"PASS {i['pass_count']}",f"FAIL {i['fail_count']}",f"N/E {i['not_evaluated_count']}",f"ERROR {i['error_count']}",i["coverage"]] for i in d["consolidated_metrics"]]
    return [_pdf_table(["Metric","Average","Evaluated","Pass","Fail","N/E","Error","Coverage"],rows,s,T,TS,c,[61*mm,22*mm,22*mm,16*mm,16*mm,16*mm,16*mm,24*mm])]
def _pdf_scenarios(d,s,T,TS,c):
    rows=[[i["scenario_id"],i["total_runs"],i["valid_runs"],i["technical_errors"],i["latest_valid_evaluation_id"] or "—",_pdf_pct(i["latest_valid_score"]),_pdf_status_label(i["latest_valid_outcome"]),_pdf_report_reference(i["latest_valid_evaluation_id"],d,s)] for i in d["scenario_results"]]
    return _workflow_pdf(d,s,True)+([_pdf_table(["Scenario","Runs","Valid","Errors","Latest Evaluation","Score","Outcome","Report"],rows,s,T,TS,c,[42*mm,14*mm,14*mm,14*mm,40*mm,18*mm,20*mm,93*mm])] if rows else [Paragraph("No evaluations match this selection.",s["BodyText"])])
def _pdf_issues(d,s):
    blocks=[]
    for kind, text, evaluation_id in _failure_content(d):
        if kind == "reference":
            blocks.extend([_pdf_report_reference(evaluation_id,d,s), Spacer(1,5)])
        else:
            blocks.append(Paragraph(_e(text), s["Subheading" if kind == "heading" else "BodyText"]))
    return blocks
def _pdf_recommendations(d,s):
    if not d["recommendations"]:return [Paragraph("No evidence-backed recommendations were generated.",s["BodyText"])]
    blocks=[]
    for i in d["recommendations"]:
        blocks.extend([Paragraph(f"<b>{_e(_pdf_status_label(i['priority']))}</b> — {_e(i['trigger'])}",s["Subheading"]),Paragraph(f"<b>Recommendation:</b> {_e(i['recommendation'])}<br/><b>Affected scenarios:</b> {_e(', '.join(i['affected_scenarios']) or '—')}",s["BodyText"])])
        blocks.extend(_pdf_report_reference(evaluation_id,d,s) for evaluation_id in i["evaluation_ids"])
        blocks.append(Spacer(1,5))
    return blocks
def _pdf_history(d,s):
    x=d["historical_comparison"]
    if x.get("status")=="comparable":return [Paragraph(f"<b>{_e(x['previous_batch'])} to {_e(x['current_batch'])}</b>: mean score change {_e(_pct(x.get('mean_score_change')))}; pass change {_e(x.get('pass_change'))}; fail change {_e(x.get('fail_change'))}.",s["BodyText"])]
    return [Paragraph("<b>Historical Comparison Unavailable</b>",s["Subheading"]),Paragraph(_e(x["message"]),s["BodyText"]),*[Paragraph(f"• {_e(r)}",s["Small"]) for r in x["reasons"]]]
def _pdf_references(d,s,T,TS,c):
    rows=[[i["scenario_id"],i["evaluation_id"],i["timestamp"] or "—",_pdf_status_label(i["outcome"]),_pdf_pct(i["score"]),i["technical_status"],_pdf_report_reference(i["evaluation_id"],d,s)] for i in d["individual_reports"]]
    return [_pdf_table(["Scenario","Evaluation ID","Timestamp","Outcome","Score","Technical","Report"],rows,s,T,TS,c,[38*mm,36*mm,47*mm,20*mm,18*mm,38*mm,64*mm])] if rows else [Paragraph("No included evaluation runs.",s["BodyText"])]


def _docx_table(document, headers, rows):
    table=document.add_table(rows=1,cols=len(headers)); table.style="Light Shading Accent 1"
    for index,value in enumerate(headers):table.rows[0].cells[index].text=str(value)
    for row in rows:
        cells=table.add_row().cells
        for index,value in enumerate(row):cells[index].text=str(value if value not in (None,"") else "—")
def _docx_executive(doc,d):
    x=d["executive_summary"];period=x.get("selected_period") or {};_docx_table(doc,["Measure","Value"],[["Mean Selected Run Score",_pct(x["mean_selected_run_score"])],["Total Evaluation Runs",x["total_evaluation_runs"]],["Unique Scenarios",x["unique_scenarios"]],["Passed / Failed / Error / Unknown",f"{x['passed_runs']} / {x['failed_runs']} / {x['technical_errors']} / {x['unknown_outcomes']}"],["Selected Evaluation Period",f"{period.get('from') or '—'} to {period.get('to') or '—'}"]]); doc.add_heading("Performance Overview",2);doc.add_paragraph(x["performance_overview"])
    doc.add_heading("Evaluation Outcome Distribution", 2)
    if not d["outcome_distribution"]["represented_total"]:
        doc.add_paragraph(_OUTCOME_EMPTY)
    _docx_table(doc, ["Outcome", "Count", "Percentage"], _outcome_rows(d))
    for note in _outcome_notes(d):
        doc.add_paragraph(note)
def _docx_coverage(doc,d):
    _workflow_docx(doc,d)
    x=d["evaluation_coverage"];doc.add_paragraph(f"Unique scenarios: {x['unique_scenarios_tested']} | Total runs: {x['total_evaluation_runs']} | Valid scored runs: {x['valid_scored_runs']} | Technical errors: {x['technical_error_runs']} | Unknown: {x['unknown_outcome_runs']}");_docx_table(doc,["Scenario","Runs","Valid","Errors","Last Evaluated"],[[i["scenario_id"],i["total_runs"],i["valid_runs"],i["technical_errors"],i["last_evaluated"]] for i in x["scenario_coverage"]]);doc.add_heading("Metric Coverage",2);_docx_table(doc,["Metric","Evaluated","Not Evaluated","Error","Coverage"],[[i["name"],i["evaluated_count"],i["not_evaluated_count"],i["error_count"],i["coverage"]] for i in x["metric_coverage"]]);[doc.add_paragraph(w,style="List Bullet") for w in x["comparability_warnings"]]
def _docx_metrics(doc,d):_docx_table(doc,["Metric","Average","Evaluated","Pass","Fail","Not Evaluated","Error","Coverage"],[[i["name"],_pct(i["average_score"]),i["evaluated_count"],i["pass_count"],i["fail_count"],i["not_evaluated_count"],i["error_count"],i["coverage"]] for i in d["consolidated_metrics"]])
def _docx_scenarios(doc,d):
    _workflow_docx(doc,d,True)
    _docx_table(doc,["Scenario","Runs","Valid","Errors","Latest Evaluation","Score","Outcome","Report Reference"],[[i["scenario_id"],i["total_runs"],i["valid_runs"],i["technical_errors"],i["latest_valid_evaluation_id"],_pct(i["latest_valid_score"]),i["latest_valid_outcome"],i["report_url"]] for i in d["scenario_results"]])
def _docx_issues(doc,d):
    for kind, text, evaluation_id in _failure_content(d):
        if kind == "heading":
            doc.add_heading(text, 2)
        else:
            doc.add_paragraph(f"Individual Report: {text}" if kind == "reference" else text)
def _docx_recommendations(doc,d):
    if not d["recommendations"]:doc.add_paragraph("No evidence-backed recommendations were generated.")
    for i in d["recommendations"]:doc.add_heading(f"{i['priority']} — {i['trigger']}",2);doc.add_paragraph(i["recommendation"]);doc.add_paragraph("Supporting evaluations: "+", ".join(i["evaluation_ids"]))
def _docx_history(doc,d):
    x=d["historical_comparison"];doc.add_paragraph(x["message"],style="Intense Quote");doc.add_paragraph(f"{x.get('previous_batch')} to {x.get('current_batch')}: mean score change {_pct(x.get('mean_score_change'))}.") if x.get("status")=="comparable" else [doc.add_paragraph(r,style="List Bullet") for r in x["reasons"]]
def _docx_references(doc,d):_docx_table(doc,["Scenario","Evaluation ID","Timestamp","Outcome","Score","Technical Status","Report Reference"],[[i["scenario_id"],i["evaluation_id"],i["timestamp"],i["outcome"],_pct(i["score"]),i["technical_status"],i["report_url"]] for i in d["individual_reports"]])
def _csv_value(value):
    if isinstance(value,dict):return "; ".join(f"{k}={v}" for k,v in value.items())
    if isinstance(value,list):return "; ".join(str(item) for item in value)
    return value
