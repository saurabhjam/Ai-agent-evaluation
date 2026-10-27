"""Outcome visualization must only project the existing selected report counts."""
from copy import deepcopy
import csv
from datetime import datetime, timezone
import io
import json
from pathlib import Path
import shutil
import subprocess

from docx import Document
from pypdf import PdfReader
import pytest
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

from app.management.consolidated_exporters import (
    _pdf_outcome_distribution, render_consolidated_csv,
    render_consolidated_docx, render_consolidated_html, render_consolidated_pdf,
)
from tests.test_management_consolidated_report import _engine, _fixture_reports, _report


def test_distribution_reuses_coverage_counts_and_preserves_source(tmp_path):
    payloads = _fixture_reports()
    before = deepcopy(payloads)
    report = _engine(tmp_path, payloads).build("chatbot:agent-1")
    distribution = report.outcome_distribution
    assert [(s.outcome, s.count) for s in distribution.slices] == [("PASS", 1), ("FAIL", 1), ("ERROR", 1)]
    assert [s.color for s in distribution.slices] == ["#22c55e", "#ef4444", "#f59e0b"]
    assert sum(s.percentage for s in distribution.slices) == pytest.approx(100)
    assert all(s.percentage == pytest.approx(100 / 3) for s in distribution.slices)
    assert distribution.represented_total == 3
    assert distribution.unknown_count == 1
    assert distribution.represented_total + distribution.unknown_count == report.executive_summary["total_evaluation_runs"]
    assert {s.outcome.lower(): s.count for s in distribution.slices} == {k:v for k,v in report.evaluation_coverage["outcome_distribution"].items() if k != "unknown"}
    assert payloads == before


@pytest.mark.parametrize("status", ["pass", "fail", "error", "unknown"])
def test_each_outcome_filter_and_zero_slices(tmp_path, status):
    report = _engine(tmp_path, _fixture_reports()).build("chatbot:agent-1", status=status)
    d = report.outcome_distribution
    assert [s.count for s in d.slices] == [int(status == k) for k in ("pass", "fail", "error")]
    assert d.represented_total == int(status != "unknown")
    assert d.unknown_count == int(status == "unknown")
    assert [s.percentage for s in d.slices] == [100 * int(status == k) for k in ("pass", "fail", "error")]
    styles = getSampleStyleSheet()
    styles.add(styles["Heading3"].clone("Subheading"))
    styles.add(styles["BodyText"].clone("Small"))
    blocks = _pdf_outcome_distribution(report.model_dump(), styles, colors)
    drawings = [b for b in blocks if isinstance(b, Drawing)]
    assert len(drawings) == int(status != "unknown")
    if drawings:
        pie = next(shape for shape in drawings[0].contents if isinstance(shape, Pie))
        assert pie.data == [1]  # No fake zero-value slices.
    pdf = render_consolidated_pdf(report)
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Evaluation Outcome Distribution" in text
    if status == "unknown":
        assert "Unknown / unclassified: 1" in text
        assert "No PASS, FAIL, or ERROR" in text


def test_all_zero_empty_state_without_misleading_chart(tmp_path):
    report = _engine(tmp_path, _fixture_reports()).build("chatbot:agent-1", scenario_id="missing")
    d = report.outcome_distribution
    assert d.represented_total == d.unknown_count == 0
    assert all(s.count == s.percentage == 0 for s in d.slices)
    html = render_consolidated_html(report)
    assert "No PASS, FAIL, or ERROR evaluation results are available for the selected report scope." in html
    assert "Unknown / unclassified:" not in html


def test_scenario_date_baseline_and_deletion_selection_stays_identical(tmp_path):
    payloads = _fixture_reports()
    engine = _engine(tmp_path, payloads)
    report = engine.build("chatbot:agent-1", scenario_id="scenario-a", date_from=datetime(2026,10,21,11,30,tzinfo=timezone.utc), date_to=datetime(2026,10,21,12,30,tzinfo=timezone.utc))
    assert report.selection.included_evaluation_ids == ["run-2"]
    assert [s.count for s in report.outcome_distribution.slices] == [1, 0, 0]
    # Store deletion removes the run from every surface, without a separate query/cache.
    payloads[:] = [p for p in payloads if p["evaluation_id"] != "run-2"]
    report = engine.build("chatbot:agent-1")
    assert [s.count for s in report.outcome_distribution.slices] == [0, 1, 1]
    assert "old" not in report.selection.included_evaluation_ids
    assert report.outcome_distribution.represented_total + report.outcome_distribution.unknown_count == len(report.individual_reports)


def test_html_docx_csv_use_the_exact_normalized_rows(tmp_path):
    report = _engine(tmp_path, _fixture_reports()).build("chatbot:agent-1")
    html = render_consolidated_html(report)
    doc = Document(io.BytesIO(render_consolidated_docx(report)))
    outcome_table = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["Outcome", "Count", "Percentage"])
    expected = [[s.outcome, str(s.count), f"{s.percentage:.1f}%"] for s in report.outcome_distribution.slices]
    assert [[c.text for c in row.cells] for row in outcome_table.rows[1:]] == expected
    for outcome, count, percentage in expected:
        assert f"<td>{outcome}</td><td>{count}</td><td>{percentage}</td>" in html
    note = "Unknown / unclassified: 1 (excluded from the pie distribution)."
    assert note in html
    assert note in "\n".join(p.text for p in doc.paragraphs)
    rows = list(csv.DictReader(io.StringIO(render_consolidated_csv(report).decode("utf-8-sig"))))
    assert [[r["name"],r["value"],r["details"]] for r in rows if r["record_type"] == "outcome_summary"] == expected


def test_pdf_uses_vector_donut_and_shared_legend_values(tmp_path):
    report = _engine(tmp_path, _fixture_reports()).build("chatbot:agent-1")
    pdf = render_consolidated_pdf(report)
    reader = PdfReader(io.BytesIO(pdf))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "Total represented: 3" in text
    assert "Unknown / unclassified: 1" in text
    assert text.count("33.3%") >= 3
    assert all(s.outcome in text for s in report.outcome_distribution.slices)
    assert any(b" c" in page.get_contents().get_data() for page in reader.pages)  # Vector curves, not a bitmap.


@pytest.mark.skipif(not shutil.which("node"), reason="Node is needed for the frontend runtime check")
@pytest.mark.parametrize("status", [None, "fail", "unknown"])
def test_web_renderer_consumes_normalized_values_without_recounting(tmp_path, status):
    report = _engine(tmp_path, _fixture_reports()).build("chatbot:agent-1", status=status)
    source = Path(__file__).resolve().parents[1] / "frontend/static/app.js"
    script = """
const fs=require('fs'), vm=require('vm');
const ctx={window:{},document:{addEventListener(){}}}; vm.createContext(ctx);
vm.runInContext(fs.readFileSync(process.argv[1],'utf8'),ctx);
const d=JSON.parse(fs.readFileSync(0,'utf8'));
process.stdout.write(ctx.window.App._consolidatedOutcomeChart(d));
"""
    html = subprocess.run([shutil.which("node"), "-e", script, str(source)], input=report.outcome_distribution.model_dump_json(), text=True, capture_output=True, check=True, timeout=20).stdout
    for s in report.outcome_distribution.slices:
        assert f"<strong>{s.outcome}</strong><span>{s.count}</span><span>{s.percentage:.1f}%</span>" in html
    if status == "unknown":
        assert "conic-gradient" not in html
        assert "excluded from the pie distribution" in html
    else:
        assert 'role="img"' in html
        assert "conic-gradient" in html
        if status == "fail":
            assert "#ef4444 0% 100%" in html
            assert "#22c55e 0%" not in html


def test_long_report_keeps_all_eight_sections_and_paginates(tmp_path):
    reports = [_report(f"run-{i}", f"scenario-{i}", "2026-10-21T12:00:00Z", passed=False, score=.2, failure=[{"metric_name":"Goal Completion", "reason":"Long existing reason. "*80}]) for i in range(12)]
    report = _engine(tmp_path, reports).build("chatbot:agent-1")
    reader = PdfReader(io.BytesIO(render_consolidated_pdf(report)))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert len(reader.pages) > 5
    assert all(title in text for title in report.section_order)
    assert "Total represented: 12" in text
